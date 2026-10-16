"""Build a plain Word proposal for discussion with the research supervisor."""
from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "documents" / "Tema_issledovaniya_dlya_nauchnogo_rukovoditelya.docx"


def font(style, size=14, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for name in ["ascii", "hAnsi", "eastAsia", "cs"]:
        fonts.set(qn("w:"+name), "Times New Roman")
    for name in ["asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        fonts.attrib.pop(qn("w:"+name), None)


def text(doc, value, style=None):
    return doc.add_paragraph(value, style=style)


def heading(doc, value):
    return doc.add_paragraph(value, style="Heading 1")


def link(paragraph, label, url):
    relationship = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    props.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "25")
    props.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    content = OxmlElement("w:t")
    content.text = label
    run.append(content)
    node.append(run)
    paragraph._p.append(node)


def math_run(value):
    node = OxmlElement("m:r")
    content = OxmlElement("m:t")
    content.text = value
    node.append(content)
    return node


def hhi_equation(doc):
    """Native editable OMML, verified in the document rendering step."""
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_after = Pt(5)
    math = OxmlElement("m:oMathPara")
    expression = OxmlElement("m:oMath")
    expression.append(math_run("HHI = "))
    nary = OxmlElement("m:nary")
    properties = OxmlElement("m:naryPr")
    char = OxmlElement("m:chr")
    char.set(qn("m:val"), "∑")
    properties.append(char)
    limits = OxmlElement("m:limLoc")
    limits.set(qn("m:val"), "undOvr")
    properties.append(limits)
    hide = OxmlElement("m:supHide")
    hide.set(qn("m:val"), "1")
    properties.append(hide)
    nary.append(properties)
    sub = OxmlElement("m:sub")
    sub.append(math_run("d"))
    nary.append(sub)
    nary.append(OxmlElement("m:sup"))
    term = OxmlElement("m:e")
    power = OxmlElement("m:sSubSup")
    base = OxmlElement("m:e")
    base.append(math_run("s"))
    power.append(base)
    index = OxmlElement("m:sub")
    index.append(math_run("d"))
    power.append(index)
    exponent = OxmlElement("m:sup")
    exponent.append(math_run("2"))
    power.append(exponent)
    term.append(power)
    nary.append(term)
    expression.append(nary)
    math.append(expression)
    paragraph._p.append(math)


def results_table(doc):
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [3.2, 3.4, 4.1, 5.0]
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    headers = ["Экспортер", "Изменение HHI", "Товарный компонент", "Географический компонент"]
    for cell, width, value in zip(table.rows[0].cells, widths, headers):
        cell.width = Cm(width)
        cell.text = value
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    with (ROOT / "results/decomposition.csv").open() as handle:
        rows = [r for r in csv.DictReader(handle) if r["metric"] == "HHI"
                and r["start"] == "2017" and r["end"] == "2024"]
    for row in rows:
        cells = table.add_row().cells
        values = ["Япония" if row["reporter"] == "392" else "Республика Корея"]
        values += [f"{float(row[k]):+.6f}".replace(".", ",").replace("-", "−")
                   for k in ["delta", "product_mix", "within_product_geography"]]
        for cell, value, width in zip(cells, values, widths):
            cell.width = Cm(width)
            cell.text = value
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        side = OxmlElement("w:"+name)
        side.set(qn("w:val"), "single")
        side.set(qn("w:sz"), "6")
        side.set(qn("w:color"), "D9D9D9")
        borders.append(side)
    table._tbl.tblPr.append(borders)
    for i, row in enumerate(table.rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                node = OxmlElement("w:"+edge)
                node.set(qn("w:w"), "80")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(12)
                    run.font.bold = i == 0
    return table


def reference(doc, number, description, urls):
    paragraph = doc.add_paragraph(style="Sources")
    paragraph.add_run(f"{number}. {description} ")
    for i, (label, url) in enumerate(urls):
        if i:
            paragraph.add_run("; ")
        link(paragraph, label, url)
    paragraph.add_run(".")


def main():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin, section.right_margin = Cm(3), Cm(2)
    section.footer_distance = Cm(1)
    normal = doc.styles["Normal"]
    font(normal)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True
    title = doc.styles["Title"]
    font(title, 14, True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.line_spacing = 1.15
    title.paragraph_format.keep_with_next = True
    h1 = doc.styles["Heading 1"]
    font(h1, 14, True)
    h1.paragraph_format.first_line_indent = Cm(0)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.space_before = Pt(10)
    h1.paragraph_format.space_after = Pt(3)
    h1.paragraph_format.line_spacing = 1.15
    h1.paragraph_format.keep_with_next = True
    sources = doc.styles.add_style("Sources", 1)
    sources.base_style = normal
    font(sources, 12.5)
    sources.paragraph_format.first_line_indent = Cm(0)
    sources.paragraph_format.line_spacing = 1.15
    sources.paragraph_format.space_after = Pt(6)
    sources.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sources.paragraph_format.keep_together = True
    note = doc.styles.add_style("Note", 1)
    note.base_style = normal
    font(note, 12)
    note.paragraph_format.first_line_indent = Cm(0)
    note.paragraph_format.line_spacing = 1.15
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(5)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    run = footer.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)
    doc.core_properties.title = "Предложение темы исследования экспорта интегральных схем"
    doc.core_properties.subject = "Материалы для обсуждения с научным руководителем"
    doc.core_properties.author = "Артём"
    doc.core_properties.keywords = "международная торговля, интегральные схемы, концентрация экспорта"
    doc.core_properties.comments = ""

    text(doc, "Предложение темы исследования экспорта интегральных схем", "Title")
    text(doc, "Материалы для обсуждения с научным руководителем\n16 октября 2026 года", "Note")
    text(doc, "Предлагаю подготовить статью о вкладе товарной структуры и перераспределения рынков сбыта в изменение концентрации экспорта интегральных схем Японии и Республики Корея. Данные собраны, предварительные результаты получены.")

    heading(doc, "1 Тема и постановка вопроса")
    p = doc.add_paragraph()
    p.add_run("Рабочее название: ").bold = True
    p.add_run("«Географическая концентрация экспорта интегральных схем Японии и Южной Кореи: роль товарной структуры и агрегации».")
    text(doc, "Тема относится к мировой экономике и международной торговле. Общая концентрация направлений зависит от товарного состава экспорта. Например, память и процессоры могут продаваться на разных рынках. Рост доли одной категории способен изменить общий индекс даже при сохранении географии поставок внутри каждой категории.")
    text(doc, "Япония и Корея выбраны для сопоставления различий товарной специализации. Доля памяти в экспорте HS8542 в 2017 году составляла около 36,5% у Японии и 70,3% у Кореи. Это различие необходимо учитывать при интерпретации общей концентрации.")

    heading(doc, "2 Цель и задачи")
    text(doc, "Цель работы — оценить вклад изменения товарной структуры и географии поставок внутри товарных категорий в изменение концентрации экспорта интегральных схем Японии и Республики Корея за 2017–2024 годы.")
    text(doc, "Для этого предполагается рассчитать показатели концентрации по годам, выполнить двухфакторное разложение их изменения, определить вклад отдельных категорий и направлений, а затем проверить чувствительность результатов к выбору периода и географической агрегации.")
    text(doc, "В исходном плане предполагалось проверить, может ли общий индекс скрывать разнонаправленные изменения двух компонентов. Знак компонентов заранее не задавался.")

    heading(doc, "3 Данные и границы исследования")
    text(doc, "Основной источник — UN Comtrade [1]. Используется собственная статистика годового экспорта Японии и Республики Корея за 2017–2024 годы. Показатель стоимости — валовой экспорт FOB в текущих долларах США. Основной расчет охватывает все представленные в статистике направления поставок; доли Китая, Гонконга, США и ЕС рассматриваются дополнительно.")
    text(doc, "Товарная граница — группа HS8542, включающая интегральные схемы и их части. Выделяются пять подкатегорий: 854231 — процессоры и контроллеры; 854232 — память; 854233 — усилители; 854239 — прочие интегральные схемы; 854290 — части. Оборудование для производства и другие полупроводниковые приборы в эту группу не входят. Названия и коды используемых подкатегорий сопоставлены между редакциями HS2017 и HS2022 [2].")
    text(doc, "В таблице содержится 4 118 наблюдений вида «экспортер — год — товар — направление». Проверены суммы поставок по направлениям и товарным подкатегориям. Повторов ключа, отрицательных и пропущенных стоимостей не обнаружено. Исходные данные и код сохранены для воспроизведения расчетов.")
    text(doc, "Китай и Гонконг в основном варианте учитываются раздельно; их объединение используется для проверки чувствительности. Для доли ЕС применяется постоянный состав ЕС-27, а в основном индексе его участники остаются отдельными направлениями. Категория Comtrade «Other Asia, nes» сохраняется под исходным названием. Доступность полного набора за 2025 год еще предстоит проверить.")

    heading(doc, "4 Метод исследования")
    text(doc, "Основной показатель — индекс Херфиндаля — Хиршмана (HHI), равный сумме квадратов долей направлений в общей стоимости экспорта:")
    hhi_equation(doc)
    text(doc, "Здесь s — доля направления d. Более высокий индекс соответствует большей концентрации. Общие доли направлений представляются через веса пяти товарных категорий и распределение поставок внутри каждой из них.")
    text(doc, "Изменение HHI раскладывается по известному принципу Шепли [3]. Рассчитываются четыре сочетания товарных весов и географических долей начального и конечного годов. Вклад каждого фактора усредняется по двум порядкам замены. Сумма товарного и географического компонентов точно равна общему изменению индекса. Это учетное разложение; оно не оценивает причинное влияние торговой политики.")
    text(doc, "Основное сравнение — 2017 и 2024 годы. Проверки включают другие базовые годы, сумму годовых разложений, объединение Китая с Гонконгом и исключение частей. Дополнительно рассчитываются доли направлений и доля трех крупнейших покупателей.")

    heading(doc, "5 Предварительные результаты")
    text(doc, "В основном варианте HHI Японии изменился с 0,152278 до 0,151452, а Кореи — с 0,269797 до 0,212029. Результаты разложения приведены ниже.")
    caption = text(doc, "Таблица 1. Разложение изменения HHI за 2017–2024 годы", "Note")
    caption.paragraph_format.keep_with_next = True
    results_table(doc)
    text(doc, "Источник: UN Comtrade, расчеты автора. Шкала HHI — 0–1.", "Note")
    text(doc, "У Японии небольшое итоговое изменение складывается из двух противоположных компонентов. У Кореи оба компонента отрицательны. При детализации географического компонента основной отрицательный вклад в обеих странах приходится на память, а процессоры и контроллеры частично его компенсируют.")
    text(doc, "В Корее географический компонент составляет 75,6% общего снижения при сравнении крайних лет и 56,3% при суммировании семи годовых разложений. Эти способы учитывают разные траектории замены факторов, сохраняя одно итоговое изменение. У Японии противоположные знаки суммарных компонентов сохраняются, но не наблюдаются в каждом годовом переходе.")
    text(doc, "Объединение Китая и Гонконга меняет снижение японского HHI с 0,000826 до 0,016460. Вывод о почти неизменной концентрации, следовательно, зависит от способа учета направлений. Эту чувствительность необходимо отразить в статье.")

    heading(doc, "6 Место в литературе и возможная новизна")
    text(doc, "В работе Hillrichs и Wölfl [4] рассматриваются концентрация торговли и различия типов полупроводников. Koh [5] анализирует корейскую отрасль в условиях торговых ограничений. Kitetu, Mbante II и Ko [6] уже сравнивают Корею и Японию по тем же пяти HS6, используя показатели торговой конкурентоспособности. Koo [7] рассчитывает концентрацию отдельных категорий корейской полупроводниковой торговли на данных до 2025 года.")
    text(doc, "Предполагаемый вклад моей работы — разложить изменение общего географического HHI, сопоставить двух экспортеров и показать чувствительность интерпретации к агрегации. В рассмотренных полных текстах прямого повторения этой постановки не обнаружено. Сам метод, пара стран и применение Comtrade новизной не являются.")
    text(doc, "У Hong [8] и Kim и Lee [9] проверены только записи, аннотации и отдельные фрагменты. Сохраняется риск совпадения с дополнительными расчетами в полных текстах. Поэтому вклад пока оценивается как возможная узкая эмпирическая добавка; его достаточность для публикации предлагаю обсудить после выбора журнала.")

    heading(doc, "7 Ограничения и дальнейшая работа")
    text(doc, "Стоимостные доли отражают одновременно цены, физические объемы и состав продукции внутри HS6. Количественные поля Японии и Кореи несопоставимы, поэтому отделение ценового эффекта в текущем варианте не проводится. Прямые таможенные направления не раскрывают конечного потребителя и реэкспорт через торговые хабы. Результаты также не распространяются автоматически на всю полупроводниковую отрасль.")
    text(doc, "Следующие шаги — уточнить отличие от ближайших исследований, проверить возможность дополнения данными за 2025 год, оформить годовую динамику и проверки чувствительности, затем подготовить текст статьи. Предполагаемая структура: введение и обзор литературы, данные и метод, результаты, проверки устойчивости, обсуждение ограничений и заключение.")
    text(doc, "Хотел бы согласовать постановку вопроса, достаточность вклада и подходящий формат публикации. Собранные данные позволяют начать подготовку первой версии статьи.")

    heading(doc, "Основные источники")
    reference(doc, 1, "United Nations. UN Comtrade Database. Годовые данные экспорта Японии и Республики Корея за 2017–2024 гг., HS8542 и пять HS6. Исходная выгрузка сохранена 15.10.2026.", [("comtradeplus.un.org", "https://comtradeplus.un.org/")])
    reference(doc, 2, "World Customs Organization. Harmonized System Nomenclature. Chapter 85. Редакции 2017 и 2022 гг.", [("HS2017", "https://www.wcoomd.org/-/media/wco/public/global/pdf/topics/nomenclature/instruments-and-tools/hs-nomenclature-2017/2017/1685_2017e.pdf?la=en"), ("HS2022", "https://www.wcoomd.org/-/media/wco/public/global/pdf/topics/nomenclature/instruments-and-tools/hs-nomenclature-2022/2022/1685_2022e.pdf?la=en")])
    reference(doc, 3, "Shorrocks A. F. Decomposition procedures for distributional analysis: a unified framework based on the Shapley value // The Journal of Economic Inequality. 2013. Vol. 11. P. 99–126.", [("DOI: 10.1007/s10888-011-9214-z", "https://doi.org/10.1007/s10888-011-9214-z")])
    reference(doc, 4, "Hillrichs D., Wölfl A. More Chips, Anyone? Demand for Semiconductors in an Era of Active Industrial Policy // EconPol Forum. 2026. Vol. 27, No. 1. P. 35–42.", [("Страница ifo", "https://www.ifo.de/en/econpol/publications/2026/article-journal/more-chips-anyone-demand-for-semiconductors-in-era-of-active-industrial-policy")])
    reference(doc, 5, "Koh W. C. Tariffs, Tech Controls, and Techno-nationalism: Korea’s Semiconductor Industry in an Era of Sanctions and Strategic Competition. AMRO Country Report ACR-2026-03. Selected Issue 1. 2026.", [("Полный текст AMRO", "https://amro-asia.org/wp-content/uploads/2026/03/Korea-ACV-2025-SI-1.pdf")])
    reference(doc, 6, "Kitetu G. M., Mbante II A. R., Ko J.-H. Revealed Trade Competitiveness between Korea and Japan: Is It Viable to Deepen Economic Integration? // Journal of Global Business and Trade. 2021. Vol. 17, No. 2. P. 67–84.", [("DOI: 10.20294/jgbt.2021.17.2.67", "https://doi.org/10.20294/jgbt.2021.17.2.67")])
    reference(doc, 7, "Koo J. Structural Changes in Korea’s Semiconductor Supply Chain under the Changing Economic Security Environment: An Analysis of Trade Structures, 2015–2025 // Public Policy Review. 2026. Vol. 40, No. 2. P. 319–339. Статья на корейском языке.", [("DOI: 10.17327/ippa.2026.40.2.011", "https://doi.org/10.17327/ippa.2026.40.2.011")])
    reference(doc, 8, "Hong M.-S. Analysis of export changes to major trading partners of the Korean semiconductor industry // The e-Business Studies. 2023. Vol. 24, No. 6. P. 179–189. Статья на корейском языке; полный текст пока не проверен.", [("DOI: 10.20462/tebs.2023.11.24.6.179", "https://doi.org/10.20462/tebs.2023.11.24.6.179")])
    reference(doc, 9, "Kim T.-H., Lee M.-J. Assessing Trade Concentration, Dependency, and Resilience in Korea’s Semiconductor Supply Chain: A Network and Entropy-Based Approach // Journal of Korea Trade. 2025. Vol. 29, No. 7. P. 67–95. Полный текст пока не проверен.", [("DOI: 10.35611/jkt.2025.29.7.67", "https://doi.org/10.35611/jkt.2025.29.7.67")])

    # Remove template borders, including Word's built-in blue Title rule.
    for root in [doc.styles.element, doc.element]:
        for border in root.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
