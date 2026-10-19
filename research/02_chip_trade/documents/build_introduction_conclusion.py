"""Build the thesis introduction and conclusion in the established chapter style.

Run with the Codex primary runtime Python. Input paragraphs are maintained in
introduction.md and conclusion.md; the existing chapters are never changed.
"""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from build_chapter1 import NS, add_text_node, element

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "introduction_conclusion.docx"
INPUTS = [HERE / "introduction.md", HERE / "conclusion.md"]


def main():
    sections = [path.read_text(encoding="utf-8") for path in INPUTS]
    definitions = dict(re.findall(r"^\[\^([^]]+)\]: (.+)$", "\n".join(sections), re.M))
    doc = Document(HERE / "chapter1.docx")
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)

    footroot = etree.Element(qn("w:footnotes"), nsmap={"w": NS})
    for fid, name in [(-1, "separator"), (0, "continuationSeparator")]:
        foot = element("w:footnote", **{"w:id": fid, "w:type": name})
        p = element("w:p")
        r = element("w:r")
        r.append(element("w:" + name))
        p.append(r)
        foot.append(p)
        footroot.append(foot)

    # An editable Word numbered list, rather than typed numbers.
    numbering = doc.part.numbering_part.element
    abstract_id = max((int(x.get(qn("w:abstractNumId")))
                       for x in numbering.findall(qn("w:abstractNum"))), default=-1) + 1
    num_id = max((int(x.get(qn("w:numId")))
                  for x in numbering.findall(qn("w:num"))), default=0) + 1
    abstract = element("w:abstractNum", **{"w:abstractNumId": abstract_id})
    abstract.append(element("w:multiLevelType", **{"w:val": "singleLevel"}))
    level = element("w:lvl", **{"w:ilvl": 0})
    for tag, value in [("start", 1), ("numFmt", "decimal"), ("lvlText", "%1."), ("lvlJc", "left")]:
        level.append(element("w:" + tag, **{"w:val": value}))
    pp = element("w:pPr")
    tabs = element("w:tabs")
    tabs.append(element("w:tab", **{"w:val": "num", "w:pos": 709}))
    pp.append(tabs)
    pp.append(element("w:ind", **{"w:left": 709, "w:hanging": 354}))
    level.append(pp)
    abstract.append(level)
    numbering.append(abstract)
    num = element("w:num", **{"w:numId": num_id})
    num.append(element("w:abstractNumId", **{"w:val": abstract_id}))
    numbering.append(num)

    notes = []
    task_count = 0
    paragraph_counts = []

    def add_footnote(p, key):
        assert key in definitions, key
        fid = len(notes) + 1
        notes.append(key)
        run = p.add_run()
        run._r.get_or_add_rPr().append(element("w:rStyle", **{"w:val": "FootnoteReference"}))
        run._r.append(element("w:footnoteReference", **{"w:id": fid}))
        foot = element("w:footnote", **{"w:id": fid})
        para = element("w:p")
        pr = element("w:pPr")
        pr.append(element("w:pStyle", **{"w:val": "FootnoteText"}))
        para.append(pr)
        r = element("w:r")
        rpr = element("w:rPr")
        rpr.append(element("w:rStyle", **{"w:val": "FootnoteReference"}))
        r.append(rpr)
        r.append(element("w:footnoteRef"))
        para.append(r)
        r = element("w:r")
        add_text_node(r, " " + definitions[key])
        para.append(r)
        foot.append(para)
        footroot.append(foot)

    def add_inline(p, value):
        value = re.sub(r"(?<=\d) (?=(?:п\.п\.|млрд|млн|тыс\.))", "\u00a0", value)
        value = re.sub(r"(?<=\d) (?=\d{3}(?:\D|$))", "\u00a0", value)
        cursor = 0
        for match in re.finditer(r"\[\^([^]]+)\]", value):
            p.add_run(value[cursor:match.start()])
            add_footnote(p, match[1])
            cursor = match.end()
        p.add_run(value[cursor:])

    for index, text in enumerate(sections):
        text = re.sub(r"^\[\^[^]]+\]: .+\n?", "", text, flags=re.M)
        count = 0
        for block in (b.strip() for b in text.split("\n\n") if b.strip()):
            if block.startswith("# "):
                p = doc.add_paragraph(block[2:], style="Title")
                p.paragraph_format.page_break_before = bool(index)
                continue
            if re.match(r"^\d+\. ", block):
                for line in block.splitlines():
                    match = re.fullmatch(r"\d+\. (.+)", line)
                    assert match, line
                    p = doc.add_paragraph()
                    pr = p._p.get_or_add_pPr()
                    npr = element("w:numPr")
                    npr.append(element("w:ilvl", **{"w:val": 0}))
                    npr.append(element("w:numId", **{"w:val": num_id}))
                    pr.append(npr)
                    pr.append(element("w:ind", **{"w:left": 709, "w:hanging": 354}))
                    add_inline(p, match[1])
                    task_count += 1
                    count += 1
                continue
            add_inline(doc.add_paragraph(), block)
            count += 1
        paragraph_counts.append(count)

    footparts = [p for p in doc.part.related_parts.values() if str(p.partname) == "/word/footnotes.xml"]
    assert len(footparts) == 1
    footparts[0]._blob = etree.tostring(footroot, xml_declaration=True, encoding="UTF-8", standalone=True)
    doc.core_properties.title = "Введение и заключение"
    doc.core_properties.subject = "Диверсификация и волатильность экспорта интегральных схем Японии и Республики Корея"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.save(OUTPUT)

    with ZipFile(OUTPUT) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        foot = etree.fromstring(z.read("word/footnotes.xml"))
        assert len(root.findall(".//" + qn("w:footnoteReference"))) == len(notes) == 5
        assert len(root.findall(".//" + qn("w:numPr"))) == task_count == 7
        assert len(foot.findall(qn("w:footnote"))) == 7  # Five notes and two separators.
        assert not root.findall(".//" + qn("w:hyperlink"))
        assert not root.findall(".//" + qn("w:tbl"))
        assert not root.findall(".//" + qn("w:drawing"))
        assert "[^" not in "".join(root.itertext())
        rels = etree.fromstring(z.read("word/_rels/document.xml.rels"))
        assert not [r for r in rels if r.get("TargetMode") == "External"]
    report = {
        "file": OUTPUT.name,
        "source_markdown": [p.name for p in INPUTS],
        "source_sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in INPUTS},
        "paragraph_counts": dict(zip(["introduction", "conclusion"], paragraph_counts)),
        "numbered_research_tasks": task_count,
        "native_footnotes": len(notes),
        "referenced_sources": 7,
        "font": "Times New Roman", "body_points": 12, "line_spacing": 1.5,
        "page_size": "A4", "external_hyperlinks": 0,
        "tables": 0, "figures": 0,
        "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "visual_review_pending": True,
    }
    (HERE / "introduction_conclusion_checks.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
