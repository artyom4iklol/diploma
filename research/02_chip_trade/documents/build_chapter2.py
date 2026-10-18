"""Build the revised chapter with native footnotes, editable math and figures."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import re,json,hashlib
from lxml import etree
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from build_chapter1 import element,font,add_text_node,NS

HERE=Path(__file__).resolve().parent
TITLE='Глава 2 Эмпирический анализ структуры и волатильности экспорта интегральных схем Японии и Республики Корея'

def mr(text,italic=False,bold=False):
    r=element('m:r');pr=element('m:rPr');pr.append(element('m:sty',**{'m:val':'bi' if bold else ('i' if italic else 'p')}));r.append(pr)
    t=element('m:t');t.text=text;r.append(t);return r

def seq(*nodes):
    out=[]
    for n in nodes:
        if isinstance(n,list):out+=n
        elif isinstance(n,str):out.append(mr(n))
        else:out.append(n)
    return out

def box(tag,nodes):
    e=element(tag)
    for n in seq(nodes):e.append(n)
    return e

def idx(base,sub=None,sup=None,bold=False):
    if sub is None and sup is None:return mr(base,True,bold)
    tag='m:sSubSup' if sub is not None and sup is not None else ('m:sSub' if sub is not None else 'm:sSup')
    e=element(tag);e.append(box('m:e',mr(base,True,bold)))
    if sub is not None:e.append(box('m:sub',mr(sub,True)))
    if sup is not None:e.append(box('m:sup',mr(sup,True)))
    return e

def frac(a,b):
    e=element('m:f');e.append(box('m:num',a));e.append(box('m:den',b));return e

def summ(sub,body):
    e=element('m:nary');pr=element('m:naryPr');pr.append(element('m:chr',**{'m:val':'∑'}));pr.append(element('m:limLoc',**{'m:val':'subSup'}));pr.append(element('m:supHide',**{'m:val':'1'}));e.append(pr)
    e.append(box('m:sub',mr(sub,True)));e.append(element('m:sup'));e.append(box('m:e',body));return e

def sigma():
    a=element('m:acc');pr=element('m:accPr');pr.append(element('m:chr',**{'m:val':'̂'}));a.append(pr);a.append(box('m:e',mr('Σ')));return a

def sqrt(body):
    a=element('m:rad');pr=element('m:radPr');pr.append(element('m:degHide',**{'m:val':'1'}));a.append(pr);a.append(element('m:deg'));a.append(box('m:e',body));return a

def inline_math(tex):
    if tex==r'\widehat{\Sigma}':return [sigma()]
    tex=tex.replace(r'\mathbf{w}','w').replace(r'\mid','|')
    nodes=[];pos=0
    pattern=re.compile(r'([A-Za-z])(?:_(\{[^}]+\}|[A-Za-z0-9]))?(?:\^(\{[^}]+\}|[A-Za-z0-9]))?')
    for m in pattern.finditer(tex):
        if m.start()>pos:nodes.append(mr(tex[pos:m.start()].replace('-','−')))
        nodes.append(idx(m[1],m[2].strip('{}') if m[2] else None,m[3].strip('{}') if m[3] else None,bold=m[1]=='w'))
        pos=m.end()
    if pos<len(tex):nodes.append(mr(tex[pos:].replace('-','−')))
    return nodes

def equation_lines(num):
    x=lambda t:idx('x','kj,'+t);X=lambda:idx('X','kj,y');V=lambda s:idx('V',s)
    if num=='2.1':return [seq(idx('g','kj,t'),' = ',frac(seq('2(',x('t'),' − ',x('t−12'),')'),seq(x('t'),' + ',x('t−12'))))]
    if num=='2.2':return [seq(idx('a','k,y'),' = ',frac(summ('j',X()),summ('k',summ('j',X()))),',     ',idx('b','j|k,y'),' = ',frac(X(),summ('j',X())))]
    if num=='2.3':return [seq(idx('G','t','rs'),' = ',summ('k',summ('j',seq(idx('w','kj','rs'),idx('g','kj,t'))))),seq(idx('w','kj','rs'),' = ',idx('a','k,yᵣ'),idx('b','j|k,yₛ'),',     r,s ∈ {0,1}')]
    if num=='2.4':
        tr=element('m:sSup');tr.append(box('m:e',seq('(',idx('w',sup='rs',bold=True),')')));tr.append(box('m:sup',mr('T')))
        return [seq(V('rs'),' = ',tr,sigma(),idx('w',sup='rs',bold=True),',     ',idx('S','rs'),' = 100',sqrt(V('rs')))]
    if num=='2.5':return [seq(idx('C','T'),' = ',frac(seq('(',V('10'),' − ',V('00'),') + (',V('11'),' − ',V('01'),')'),'2')),seq(idx('C','G'),' = ',frac(seq('(',V('01'),' − ',V('00'),') + (',V('11'),' − ',V('10'),')'),'2'))]
    if num=='2.6':
        def sw(s):
            a=element('m:sSub');a.append(box('m:e',seq('(',sigma(),idx('w',sup=s,bold=True),')')));a.append(box('m:sub',mr('i',True)));return a
        return [seq(idx('d','i','A→B'),' = (',idx('w','i','B'),' − ',idx('w','i','A'),')[',sw('B'),' + ',sw('A'),']'),seq(summ('i',idx('d','i','A→B')),' = ',V('B'),' − ',V('A'))]
    raise ValueError(num)

def main():
    parts=[(HERE/f'chapter2_{i}.md').read_text().strip() for i in range(1,5)]
    main_text='# '+TITLE+'\n\n'+'\n\n'.join(parts)+'\n'
    (HERE/'chapter2_content.md').write_text(main_text)
    appendix=(HERE/'chapter2_appendices.md').read_text()
    audit=json.loads((HERE/'chapter2_source_audit.json').read_text())
    bibliography=['# Список использованных источников ко второй главе','']
    for i,s in enumerate(sorted(audit['sources'],key=lambda s:s['citation'].casefold()),1):
        line=f'{i}. '+s['citation']
        consulted=s.get('consulted_citation')
        if consulted and consulted!=s['citation']:line+=' Использованная версия: '+consulted
        if s['source_id']=='shorrocks_2013':line=f'{i}. '+s['consulted_citation']
        bibliography += [line,'']
    text=main_text+'\n'+appendix+'\n'+'\n'.join(bibliography)
    defs=dict(re.findall(r'^\[\^([^]]+)\]: (.+)$',text,re.M));text=re.sub(r'^\[\^[^]]+\]: .+\n?', '',text, flags=re.M)
    doc=Document(HERE/'chapter1.docx')
    for child in list(doc.element.body):
        if child.tag!=qn('w:sectPr'):doc.element.body.remove(child)
    for style in doc.styles:
        if hasattr(style,'font') and style.name not in ['Title','Heading 1','Footnote Text','Footnote Reference']:font(style.font)
    footroot=etree.Element(qn('w:footnotes'),nsmap={'w':NS})
    for fid,name in [(-1,'separator'),(0,'continuationSeparator')]:
        f=element('w:footnote',**{'w:id':fid,'w:type':name});p=element('w:p');r=element('w:r');r.append(element('w:'+name));p.append(r);f.append(p);footroot.append(f)
    notes=[];math_count=0;image_count=0;table_count=0

    def note(p,key):
        assert key in defs;fid=len(notes)+1;notes.append(key)
        run=p.add_run();rp=run._r.get_or_add_rPr();rp.append(element('w:rStyle',**{'w:val':'FootnoteReference'}));run._r.append(element('w:footnoteReference',**{'w:id':fid}))
        f=element('w:footnote',**{'w:id':fid});para=element('w:p');pr=element('w:pPr');pr.append(element('w:pStyle',**{'w:val':'FootnoteText'}));para.append(pr)
        r=element('w:r');rp=element('w:rPr');rp.append(element('w:rStyle',**{'w:val':'FootnoteReference'}));r.append(rp);r.append(element('w:footnoteRef'));para.append(r)
        r=element('w:r');add_text_node(r,' '+defs[key]);para.append(r);f.append(para);footroot.append(f)

    def inline(p,content):
        content=re.sub(r'(?<=\d) (?=\d{3}(?:\D|$))','\u00a0',content)
        content=re.sub(r'(?<=\d) (?=(?:п\.п\.|млрд|млн|тыс\.))','\u00a0',content)
        pos=0
        for m in re.finditer(r'\[\^([^]]+)\]|\$([^$]+)\$',content):
            p.add_run(content[pos:m.start()])
            if m[1]:note(p,m[1])
            else:
                math=element('m:oMath')
                for n in inline_math(m[2]):math.append(n)
                p._p.append(math)
            pos=m.end()
        p.add_run(content[pos:])

    def normal_p(content,size=None):
        p=doc.add_paragraph();inline(p,content)
        if size:
            for r in p.runs:font(r.font,size)
        return p

    blocks=[b.strip() for b in text.split('\n\n') if b.strip()]
    in_bib=False;in_appendix=False;first_appendix=True;last_caption=''
    for b in blocks:
        if b.startswith('# '):
            label=b[2:]
            if label==TITLE:p=doc.add_paragraph(label,style='Title')
            else:
                p=doc.add_paragraph(label,style='Heading 1');p.paragraph_format.page_break_before=True
                in_appendix=label.startswith('Приложения');in_bib=label.startswith('Список')
            continue
        if b.startswith('## '):
            p=doc.add_paragraph(b[3:],style='Heading 1')
            if in_appendix:
                p.paragraph_format.page_break_before=not first_appendix
                first_appendix=False
            continue
        if b.startswith('$$'):
            number=re.search(r'\\tag\{([^}]+)\}',b)[1];p=doc.add_paragraph();pf=p.paragraph_format
            pf.first_line_indent=Cm(0);pf.line_spacing=1.1;pf.space_before=Pt(6);pf.space_after=Pt(6);pf.keep_together=True
            pf.tab_stops.add_tab_stop(Cm(7.6),WD_TAB_ALIGNMENT.CENTER);pf.tab_stops.add_tab_stop(Cm(16.2),WD_TAB_ALIGNMENT.RIGHT)
            p.add_run('\t');math=element('m:oMath');lines=equation_lines(number)
            if len(lines)==1:
                for n in lines[0]:math.append(n)
            else:
                arr=element('m:eqArr');pr=element('m:eqArrPr');pr.append(element('m:baseJc',**{'m:val':'center'}));arr.append(pr)
                for ln in lines:arr.append(box('m:e',ln))
                math.append(arr)
            p._p.append(math);p.add_run('\t('+number+')');math_count+=1
            continue
        if b.startswith('!['):
            match=re.fullmatch(r'!\[([^]]+)\]\(([^)]+)\)',b);assert match
            p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(6);p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            sh=p.add_run().add_picture(str(HERE/match[2]),width=Cm(16.2));sh._inline.docPr.set('descr',match[1]);image_count+=1
            continue
        if b.startswith('**Таблица') or b.startswith('**Рисунок'):
            label=b.strip('*');last_caption=label;p=normal_p(label,11);pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.line_spacing=1.05;pf.space_before=Pt(5);pf.space_after=Pt(5);pf.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if label.startswith('Рисунок') else WD_ALIGN_PARAGRAPH.LEFT
            continue
        if b.startswith('|'):
            rows=[[v.strip() for v in line.strip('|').split('|')] for line in b.splitlines()];rows=[rows[0],*rows[2:]];n=len(rows[0]);assert all(len(row)==n for row in rows)
            widths={2:[5.3,11.2],3:[6,5.25,5.25],4:[6.6,1.6,4.15,4.15],5:[6.5,2.5,2.5,2.5,2.5],6:[4.4,1.2,2.7,2.7,2.7,2.8]}[n]
            if last_caption.startswith('Таблица 2.2'):widths=[2.4,3.4,3.5,7.2]
            table=doc.add_table(rows=0,cols=n);table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
            for c,w in zip(table.columns,widths):c.width=Cm(w)
            props=table._tbl.tblPr; borders=element('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:borders.append(element('w:'+edge,**{'w:val':'single','w:sz':5,'w:color':'D9D9D9'}))
            props.append(borders)
            for ri,row in enumerate(rows):
                cells=table.add_row().cells;trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(element('w:cantSplit'))
                if ri==0:trpr.append(element('w:tblHeader'))
                for ci,(cell,val,w) in enumerate(zip(cells,row,widths)):
                    cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcpr=cell._tc.get_or_add_tcPr();marg=element('w:tcMar')
                    for edge,v in [('top',40 if in_appendix else 65),('bottom',40 if in_appendix else 65),('left',75),('right',75)]:marg.append(element('w:'+edge,**{'w:w':v,'w:type':'dxa'}))
                    tcpr.append(marg)
                    if ri==0:tcpr.append(element('w:shd',**{'w:fill':'EDEDED'}))
                    p=cell.paragraphs[0];pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.line_spacing=1.0 if in_appendix else 1.05;pf.space_before=pf.space_after=Pt(0);pf.keep_with_next=(not in_appendix and len(rows)<=12) or ri==len(rows)-1;pf.keep_together=True
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 or (n==2) else WD_ALIGN_PARAGRAPH.CENTER
                    r=p.add_run(val);font(r.font,10.5,bold=ri==0)
            table_count+=1
            continue
        if b.startswith(('Примечание:','Источник:')):
            p=normal_p(b,10);pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.line_spacing=1;pf.space_before=Pt(3);pf.space_after=Pt(4);pf.keep_with_next=b.startswith('Примечание:')
            continue
        p=normal_p(b)
        if in_bib:
            pf=p.paragraph_format;pf.line_spacing=1.15;pf.space_after=Pt(5);pf.first_line_indent=Cm(-.7);pf.left_indent=Cm(.7)

    related=[part for rel,part in doc.part.related_parts.items() if part.partname=='/word/footnotes.xml'];assert len(related)==1
    related[0]._blob=etree.tostring(footroot,xml_declaration=True,encoding='UTF-8',standalone=True)
    doc.core_properties.title=TITLE;doc.core_properties.subject='Япония и Республика Корея, 2017–2024';doc.core_properties.author='';doc.core_properties.comments=''
    path=HERE/'chapter2.docx';doc.save(path)
    with ZipFile(path) as z:
        root=etree.fromstring(z.read('word/document.xml'));assert not root.findall('.//'+qn('w:hyperlink'));assert len(root.findall('.//'+qn('w:drawing')))==8
        assert len(root.findall('.//'+qn('w:footnoteReference')))==len(notes)==11
        assert len(root.findall('.//'+qn('w:tbl')))==table_count==11
        assert b'$$' not in z.read('word/document.xml') and b'\\tag' not in z.read('word/document.xml')
    report={'file':path.name,'native_footnotes':len(notes),'numbered_editable_equations':math_count,'all_native_math_objects':len(root.findall('.//'+qn('m:oMath'))),'figures':image_count,'main_tables':5,'appendix_tables':table_count-5,'font':'Times New Roman','body_points':12,'body_line_spacing':1.5,'page_size':'A4','external_hyperlinks':0,'source_count':len(audit['sources']),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'visual_review_pending':True}
    (HERE/'chapter2_document_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':main()
