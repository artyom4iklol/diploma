"""Insert the reviewed 30 figures into the assembled thesis (without rerunning models).

Run after build_thesis.py, or pass --base an untouched assembled thesis. Assets
are the reviewed PNG directory; 28/30 use cropped columns for print legibility.
The source text, equations, footnotes, tables and bibliography are preserved.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import argparse
import hashlib
import json
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image

HERE = Path(__file__).resolve().parent

# Source paragraph indices refer to the assembled 88-page version of 20 October.
# They are validated below; the script fails rather than editing another version.
FIGURES = [
 (1,'1.1',53,'Международная цепочка создания интегральных схем','Основные стадии международной цепочки создания стоимости и их взаимосвязи представлены на рисунке 1.1.'),
 (2,'1.2',86,'Структура экспорта и условная волатильность','Связь товарных долей, географии поставок и совместных колебаний схематически показана на рисунке 1.2. Количественная реализация этой схемы рассматривается во второй главе.'),
 (3,'1.3',93,'Число рынков и корреляция экспортных потоков','Условный пример на рисунке 1.3 показывает, почему увеличение числа одинаково взвешенных рынков дает разный результат при разных корреляциях потоков. Это иллюстрация формулы дисперсии, а не оценка по странам исследования.'),
 (4,'2.1',174,'Месячная стоимость экспорта интегральных схем и годовые темпы ее изменения',''),
 (5,'2.2',185,'Индексы стоимости экспорта четырех товарных групп, 2017 = 100','Различия в траекториях четырех товарных групп сопоставлены на рисунке 2.2. Индексы построены по годовой номинальной стоимости; логарифмическая шкала позволяет сравнить относительные изменения групп разного размера.'),
 (6,'2.3',187,'Товарная структура экспорта интегральных схем',''),
 (7,'2.4',191,'Товарный вклад в изменение номинальной стоимости экспорта между 2017 и 2024 годами','Абсолютное изменение экспорта между начальным и конечным годами разложено по товарным группам на рисунке 2.4. Это стоимостное разложение; изменение цен и физических объемов в нем не разделяется.'),
 (8,'2.5',191,'Товарная и географическая концентрация экспорта','Рисунок 2.5 сопоставляет динамику товарного и географического HHI. Географический показатель рассчитан по всем исходным направлениям, поэтому его не следует отождествлять с концентрацией на укрупненной сетке основной модели.'),
 (9,'2.6',194,'Изменение долей направлений в экспорте интегральных схем',''),
 (10,'2.7',200,'Географическая структура экспорта внутри товарных групп','Рисунок 2.7 показывает распределение каждой товарной группы по направлениям в двух опорных годах. Доли внутри группы дополняют агрегированное сравнение географии и позволяют увидеть неоднородность поставок.'),
 (11,'2.8',207,'Динамика экспорта в двух сопоставляемых эпизодах','На рисунке 2.8 квартальные стоимости нормированы к четвертому кварталу года, предшествующего соответствующему эпизоду. Сопоставление описывает наблюдавшиеся траектории, без сезонной корректировки и без выделения причинного эффекта событий.'),
 (12,'2.9',216,'Стандартное отклонение четырех условных структур экспорта',''),
 (13,'2.10',225,'Товарный и географический вклады в изменение условной дисперсии',''),
 (14,'2.11',236,'Доля памяти и состав продукции вне памяти в товарном вкладе',''),
 (15,'2.12',244,'Собственные дисперсии и ковариации в географическом вкладе',''),
 (16,'2.13',254,'Детализация географического вклада по направлениям','Детальные оценки и их интервалы сопоставлены на рисунке 2.13. Это поисковая атрибуция без поправки на множественные сравнения; вклад направления включает его ковариации с остальными потоками.'),
 (17,'2.14',273,'Чувствительность изменения стандартного отклонения к исключению одного года',''),
 (18,'3.1',294,'Сдвиг доли памяти и изменение условного стандартного отклонения','Для сопоставления направлений структурного изменения на рисунке 3.1 приведены сценарии сдвига доли памяти относительно структуры 2024 года. География внутри товаров и относительный состав остальных категорий фиксированы, история темпов общая — 2018–2024 годы; показаны поточечные 95%-е интервалы. Методика и границы интерпретации подробнее рассматриваются в разделе 3.4.'),
 (19,'3.2',296,'Необходимое расширение экспорта продукции вне памяти','Рисунок 3.2 показывает арифметический масштаб снижения доли памяти за счет расширения остальных поставок при неизменной стоимости памяти. Дополнительная стоимость экспорта не является инвестиционными расходами, бюджетом поддержки или прогнозом выручки.'),
 (20,'3.3',339,'Сопоставление географических сценариев с пропорциональным расширением','Сравнение направлений и одновременные 95%-е интервалы представлены на рисунке 3.3. Нулем служит размещение той же дополнительной товарной корзины пропорционально исходной географии, а не отсутствие прироста экспорта.'),
 (21,'3.4',343,'Требуемое расширение ячеек «товар — рынок»','Максимальный относительный прирост отдельной товарной ячейки в каждом сценарии показан на рисунке 3.4. Такая детализация необходима для оценки масштаба предполагаемых новых заказов.'),
 (22,'3.5',383,'Этапы отбора, проверки и реализации экспортного проекта',''),
 (24,'Б.1',461,'Корреляции годовых темпов роста направлений экспорта','Для каждого экспортера рисунок Б.1 сопоставляет корреляции агрегированных направлений. Это описательное представление связей; матрица основной модели рассчитана на более детальной сетке из 36 товарно-географических ячеек.'),
 (26,'Б.2',461,'Поставки от товарных групп к направлениям экспорта в 2024 году','Рисунок Б.2 связывает товарный состав и непосредственные направления зарегистрированного экспорта. Ширина потоков соответствует стоимости; схема не восстанавливает дальнейший реэкспорт, конечное потребление или весь производственный маршрут.'),
 (27,'В.1',468,'Основные оценки и интервалы при бутстреп-блоках 6 и 12 месяцев','Интервалы при двух длинах блока сопоставлены на рисунке В.1. Разность стандартных отклонений измеряется в процентных пунктах, вклады в дисперсию — в квадратных процентных пунктах; их масштабы не объединяются.'),
 (28,'В.2',468,'Основная оценка и вклады на двух исторических окнах','На рисунке В.2 веса 2017 и 2024 годов одинаковы для обоих окон, а история совместных колебаний различается. Для окон представлены точечные оценки; отдельные интервалы не рассчитывались.'),
 (23,'Д.1',479,'Месячные годовые темпы роста экспорта по товарным группам','Рисунок Д.1 показывает симметричные годовые темпы роста агрегированных товарных групп. Их не следует отождествлять с фиксированными индексами отдельных товарно-географических ячеек.'),
 (25,'Д.2',479,'Временные ряды условных индексов при структурах 2017 и 2024 годов','На рисунке Д.2 две структуры сопоставлены на одних и тех же 84 месячных наблюдениях 2018–2024 годов. Это условные исторические индексы с фиксированными весами, а не фактические агрегированные темпы экспорта.'),
 (30,'Е.1',485,'Сверка сумм месячных и годовых данных','Рисунок Е.1 иллюстрирует сверку 96 пар «экспортер — год — код». Помимо четырех основных категорий, контроль включает итог HS8542 и части HS854290; это расширенный контроль исходных данных, а не изменение товарного охвата основной модели.'),
 (29,'Ж.1',492,'Чувствительность сценариев доли памяти к историческому окну','Рисунок Ж.1 сопоставляет сценарные разности на полной истории и на двух окнах при неизменных весах. Различия линий характеризуют зависимость результата от истории темпов и не являются прогнозом.'),
]

OLD_BLOCKS = [range(i,i+4) for i in [175,188,195,217,226,237,245,274,345,368,384]]
REF_MAP = {'2.1':'2.1','2.2':'2.3','2.3':'2.6','2.4':'2.9','2.5':'2.10',
           '2.6':'2.11','2.7':'2.12','2.8':'2.14','3.1':'3.5','3.2':'3.1','3.3':'3.5'}

def sha(data): return hashlib.sha256(data).hexdigest()

def replace_refs(paragraph):
    # Replace only w:t ranges; native equations/footnote references stay untouched.
    nodes=paragraph._p.findall('.//'+qn('w:t'))
    full=''.join(n.text or '' for n in nodes)
    matches=list(re.finditer(r'(рисун(?:ок|к[а-я]*)\s+)([23]\.\d+)',full,re.I))
    for m in reversed(matches):
        old=m.group(2); new=REF_MAP.get(old,old)
        if old==new: continue
        left,right=m.span(2); offset=0
        for node in nodes:
            s=node.text or ''; end=offset+len(s)
            if end>left and offset<right:
                a=max(0,left-offset);b=min(len(s),right-offset)
                node.text=s[:a]+(new if offset<=left else '')+s[b:]
            offset=end

def add_after(anchor,text='',kind='body'):
    node=OxmlElement('w:p');anchor._p.addnext(node)
    p=Paragraph(node,anchor._parent);p.style='Normal'
    pf=p.paragraph_format
    if kind=='body':
        pf.first_line_indent=Cm(1.25);pf.line_spacing=1.5
        pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.space_after=Pt(0)
    else:
        pf.first_line_indent=Cm(0);pf.line_spacing=1
        pf.space_before=Pt(4 if kind=='image' else 2)
        pf.space_after=Pt(5 if kind=='source' else 2)
        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if kind in ['image','caption'] else WD_ALIGN_PARAGRAPH.LEFT
    pf.keep_together=True
    pf.keep_with_next=kind in ['image','caption']
    if text:
        r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(14 if kind=='body' else 11 if kind=='caption' else 10)
        r.font.color.rgb=RGBColor(0,0,0)
    return p

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--base',type=Path,required=True)
    ap.add_argument('--assets',type=Path,required=True);ap.add_argument('--output',type=Path,default=HERE/'diploma_full.docx')
    args=ap.parse_args();doc=Document(args.base);original=list(doc.paragraphs)
    assert len(original)==500 and len(doc.inline_shapes)==11, 'Unexpected base; rebuild untouched thesis first.'
    assert original[275].text.startswith('Рисунок 2.8') and original[385].text.startswith('Рисунок 3.3')
    with ZipFile(args.base) as z:
        notes=z.read('word/footnotes.xml');equations=z.read('word/document.xml').count(b'<m:oMath>')
    removed={i for block in OLD_BLOCKS for i in block}
    for i,p in enumerate(original):
        if i not in removed:replace_refs(p)
    for i in sorted(removed,reverse=True):
        node=original[i]._p;node.getparent().remove(node)
    tails={};manifest=[]
    for number,label,index,title,intro in FIGURES:
        anchor=tails.get(index,original[index])
        # 7/8 attach after the new 2.3 block, whose old source paragraph is removed.
        if index==191:anchor=tails.get(187,original[187]) if index not in tails else tails[index]
        if intro:anchor=add_after(anchor,intro)
        files=sorted(args.assets.glob(f'{number:02d}_*.png'))
        if number in (28,30):files=[p for p in files if '_part' in p.stem]
        else:files=[p for p in files if '_part' not in p.stem]
        assert files and (len(files)==2 if number in (28,30) else len(files)==1), (number,files)
        entry={'archive_number':number,'number':label,'caption':title,'source_anchor':index,'images':[]}
        for j,path in enumerate(files):
            anchor=add_after(anchor,kind='image')
            w,h=Image.open(path).size;width=min(16.5,21.5*w/h)
            shape=anchor.add_run().add_picture(str(path),width=Cm(width))
            shape._inline.docPr.set('descr',f'Рисунок {label}. {title}'+(' (продолжение)' if j else ''))
            cap=f'Рисунок {label} — {title}' if j==0 else f'Продолжение рисунка {label}'
            anchor=add_after(anchor,cap,'caption')
            source='Источник: составлено автором.' if number in (1,2,3,22) else 'Источник: расчеты автора по данным UN Comtrade.'
            anchor=add_after(anchor,source,'source')
            entry['images'].append({'file':path.name,'sha256':sha(path.read_bytes()),'width_cm':width})
        tails[index]=anchor;manifest.append(entry)
    # Remove old unreferenced image relationships (the new drawing references stay).
    used={e.get(qn('r:embed')) for e in doc.element.body.findall('.//'+qn('a:blip'))}
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in used:del doc.part.rels[rid]
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    with ZipFile(args.output) as z:
        assert z.read('word/footnotes.xml')==notes,'Footnotes changed'
        assert z.read('word/document.xml').count(b'<m:oMath>')==equations,'Equations changed'
    final=Document(args.output)
    assert len(final.tables)==23 and len(final.inline_shapes)==32
    captions=[p.text for p in final.paragraphs if re.match(r'^Рисунок [123БВДЕЖ]\.\d+ — ',p.text)]
    assert len(captions)==30,len(captions)
    report={'base_sha256':sha(args.base.read_bytes()),'output_sha256':sha(args.output.read_bytes()),
            'figures':manifest,'logical_figures':30,'image_panels':32,'preserved_tables':23,
            'preserved_equations':equations,'preserved_footnotes_sha256':sha(notes),
            'corrections':['18: explicit country headings','20: four decimal places instead of rounded zeros',
                           '25: common vertical range expanded to include -67.567265 pp'],
            'layout_adaptation':'28 and 30 split into column panels, same logical figure number'}
    (HERE/'figure_insertion_manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k!='figures'},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
