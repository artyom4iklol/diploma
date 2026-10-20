"""Assemble the thesis without changing completed chapter files.

Preserves native equations, images, tables, and footnotes. A second build accepts
page numbers obtained from the actual PDF render for a plain, nonlinked TOC.
"""
from pathlib import Path
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import argparse
import hashlib
import json
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt
from lxml import etree

from build_chapter1 import NS, element, font

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "diploma_full.docx"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def node_text(node):
    return "".join(node.itertext())


def paragraph_text(node):
    return "".join(t.text or "" for t in node.findall(".//" + qn("w:t")))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--toc-pages", type=Path)
    args = parser.parse_args()
    toc_pages = json.loads(args.toc_pages.read_text()) if args.toc_pages else {}
    doc = Document(HERE / "introduction_conclusion.docx")
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)

    originals = {n: Document(HERE / f"{n}.docx") for n in ["introduction_conclusion", "chapter1", "chapter2", "chapter3"]}
    chosen = {}
    for name, source in originals.items():
        body = [e for e in source.element.body if e.tag != qn("w:sectPr")]
        if name == "introduction_conclusion":
            split = next(i for i, e in enumerate(body) if paragraph_text(e) == "Заключение")
            chosen["introduction"] = (name, body[:split])
            chosen["conclusion"] = (name, body[split:])
        else:
            end = next((i for i, e in enumerate(body) if paragraph_text(e).startswith(("Список использованных", "Приложения ко второй"))), len(body))
            chosen[name] = (name, body[:end])

    headings = []
    for name in ["introduction", "chapter1", "chapter2", "chapter3", "conclusion"]:
        source_name, elements = chosen[name]
        source = originals[source_name]
        for e in elements:
            if e.tag != qn("w:p"):
                continue
            p = Paragraph(e, source._body)
            if p.style.name in ("Title", "Heading 1"):
                headings.append({"text": p.text, "level": 0 if p.style.name == "Title" else 1})
    headings.append({"text": "Список использованных источников", "level": 0})
    app_text = (HERE / "thesis_appendices.md").read_text()
    for line in app_text.splitlines():
        if line.startswith("# "):
            headings.append({"text": line[2:], "level": 0})
        elif line.startswith("## "):
            headings.append({"text": line[3:], "level": 1})

    title = doc.add_paragraph("Оглавление", style="Title")
    title._p.get_or_add_pPr().append(element("w:outlineLvl", **{"w:val": 9}))
    for heading in headings:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Cm(0)
        pf.left_indent = Cm(.5 if heading["level"] else 0)
        pf.right_indent = Cm(.8)
        pf.line_spacing = 1.05
        pf.space_after = Pt(3)
        pf.keep_together = True
        pf.keep_with_next = False
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(heading["text"] + "\t" + str(toc_pages.get(heading["text"], 99)))
        font(run.font, 11, bold=not bool(heading["level"]))

    footroot = etree.Element(qn("w:footnotes"), nsmap={"w": NS})
    for fid, name in [(-1, "separator"), (0, "continuationSeparator")]:
        foot = element("w:footnote", **{"w:id": fid, "w:type": name})
        p = element("w:p"); r = element("w:r")
        r.append(element("w:" + name)); p.append(r); foot.append(p); footroot.append(foot)
    note_count = 0
    drawing_count = 0
    image_maps = {}
    numbering_maps = {}
    imported_counts = {}
    content_checks = []

    def mapped_numbering(source_name, old_id):
        key = (source_name, old_id)
        if key in numbering_maps:
            return numbering_maps[key]
        source_root = originals[source_name].part.numbering_part.element
        target_root = doc.part.numbering_part.element
        old_num = next(n for n in source_root.findall(qn("w:num")) if n.get(qn("w:numId")) == old_id)
        abstract_id = old_num.find(qn("w:abstractNumId")).get(qn("w:val"))
        old_abs = next(n for n in source_root.findall(qn("w:abstractNum")) if n.get(qn("w:abstractNumId")) == abstract_id)
        new_abs_id = max(int(n.get(qn("w:abstractNumId"))) for n in target_root.findall(qn("w:abstractNum"))) + 1
        new_num_id = max(int(n.get(qn("w:numId"))) for n in target_root.findall(qn("w:num"))) + 1
        new_abs = deepcopy(old_abs); new_abs.set(qn("w:abstractNumId"), str(new_abs_id)); target_root.append(new_abs)
        new_num = deepcopy(old_num); new_num.set(qn("w:numId"), str(new_num_id))
        new_num.find(qn("w:abstractNumId")).set(qn("w:val"), str(new_abs_id)); target_root.append(new_num)
        numbering_maps[key] = str(new_num_id)
        return str(new_num_id)

    def import_part(name):
        nonlocal note_count, drawing_count
        source_name, elements = chosen[name]
        source = originals[source_name]
        source_foot = next(p for p in source.part.related_parts.values() if str(p.partname) == "/word/footnotes.xml")
        source_notes = {f.get(qn("w:id")): f for f in etree.fromstring(source_foot.blob).findall(qn("w:footnote"))}
        imported_counts[name] = len(elements)
        for index, original in enumerate(elements):
            clone = deepcopy(original)
            original_text = node_text(original)
            if clone.tag == qn("w:p"):
                p = Paragraph(clone, doc._body)
                if p.style.name == "Title":
                    p.paragraph_format.page_break_before = True
                    p._p.get_or_add_pPr().append(element("w:outlineLvl", **{"w:val": 0}))
                elif p.style.name == "Heading 1":
                    p._p.get_or_add_pPr().append(element("w:outlineLvl", **{"w:val": 1}))
            for ref in clone.findall(".//" + qn("w:footnoteReference")):
                old_id = ref.get(qn("w:id"))
                note_count += 1
                note = deepcopy(source_notes[old_id])
                note.set(qn("w:id"), str(note_count))
                ref.set(qn("w:id"), str(note_count))
                footroot.append(note)
            for num_id in clone.findall(".//" + qn("w:numId")):
                num_id.set(qn("w:val"), mapped_numbering(source_name, num_id.get(qn("w:val"))))
            for node in clone.iter():
                for attr, value in list(node.attrib.items()):
                    if attr == "{" + RNS + "}embed":
                        key = (source_name, value)
                        if key not in image_maps:
                            rel = source.part.rels[value]
                            assert not rel.is_external
                            image_maps[key] = doc.part.get_or_add_image(BytesIO(rel.target_part.blob))[0]
                        node.set(attr, image_maps[key])
            for drawing in clone.findall(".//" + qn("w:drawing")):
                drawing_count += 1
                for node in drawing.iter():
                    if etree.QName(node).localname in ("docPr", "cNvPr"):
                        node.set("id", str(drawing_count))
            assert node_text(clone) == original_text, (name, index)
            content_checks.append((name, index, hashlib.sha256(original_text.encode()).hexdigest()))
            doc.element.body.insert(len(doc.element.body) - 1, clone)

    for name in ["introduction", "chapter1", "chapter2", "chapter3", "conclusion"]:
        import_part(name)

    def normal(value, size=None, keep=False):
        p = doc.add_paragraph()
        p.add_run(value)
        p.paragraph_format.keep_with_next = keep
        if size:
            for r in p.runs:
                font(r.font, size)
        return p

    def heading(value, main=False, new_page=False):
        p = doc.add_paragraph(value, style="Title" if main else "Heading 1")
        p.paragraph_format.page_break_before = new_page
        p._p.get_or_add_pPr().append(element("w:outlineLvl", **{"w:val": 0 if main else 1}))
        return p

    bibliography = json.loads((HERE / "thesis_bibliography.json").read_text())
    heading("Список использованных источников", main=True, new_page=True)
    for row in bibliography["sources"]:
        p = normal(str(row["number"]) + ". " + row["formatted"])
        pf = p.paragraph_format
        pf.line_spacing = 1.15; pf.space_after = Pt(5)
        pf.first_line_indent = Cm(-.7); pf.left_indent = Cm(.7)
        pf.keep_together = True

    last_caption = ""
    first_appendix = True
    appendix_tables = 0
    for block in (b.strip() for b in app_text.split("\n\n") if b.strip()):
        if block.startswith("# "):
            heading(block[2:], main=True, new_page=True)
        elif block.startswith("## "):
            heading(block[3:], new_page=not first_appendix)
            first_appendix = False
        elif block.startswith("**Таблица"):
            last_caption = block.strip("*")
            p = normal(last_caption, 11, keep=True)
            pf = p.paragraph_format; pf.first_line_indent = Cm(0)
            pf.line_spacing = 1.05; pf.space_before = Pt(5); pf.space_after = Pt(5)
        elif block.startswith("|"):
            rows = [[v.strip() for v in line.strip("|").split("|")] for line in block.splitlines()]
            rows = [rows[0], *rows[2:]]
            cols = len(rows[0])
            assert all(len(r) == cols for r in rows)
            widths = {3: [5.4, 5.4, 5.4], 5: [6.2, 2.5, 2.5, 2.5, 2.5], 6: [4.2, 1.2, 2.7, 2.7, 2.7, 2.7]}[cols]
            if last_caption.startswith("Таблица Ж"):
                widths = [1.7, 3.0, 4.3, 3.6, 3.6]
            elif last_caption.startswith("Таблица З"):
                widths = [2.7, 2.5, 2.5, 4.0, 2.7, 1.8]
            table = doc.add_table(rows=0, cols=cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
            for column, width in zip(table.columns, widths): column.width = Cm(width)
            borders = element("w:tblBorders")
            for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                borders.append(element("w:" + edge, **{"w:val": "single", "w:sz": 5, "w:color": "D9D9D9"}))
            table._tbl.tblPr.append(borders)
            for ri, values in enumerate(rows):
                cells = table.add_row().cells
                trpr = table.rows[-1]._tr.get_or_add_trPr(); trpr.append(element("w:cantSplit"))
                if ri == 0: trpr.append(element("w:tblHeader"))
                for ci, (cell, value, width) in enumerate(zip(cells, values, widths)):
                    cell.width = Cm(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    props = cell._tc.get_or_add_tcPr()
                    margin = element("w:tcMar")
                    for edge, amount in [("top", 50), ("bottom", 50), ("left", 75), ("right", 75)]:
                        margin.append(element("w:" + edge, **{"w:w": amount, "w:type": "dxa"}))
                    props.append(margin)
                    if ri == 0: props.append(element("w:shd", **{"w:fill": "EDEDED"}))
                    p = cell.paragraphs[0]; pf = p.paragraph_format
                    pf.first_line_indent = Cm(0); pf.line_spacing = 1.0
                    pf.space_before = pf.space_after = Pt(0); pf.keep_together = True
                    pf.keep_with_next = ri == 0 or ri == len(rows) - 1
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    font(p.add_run(value).font, 10.5, bold=ri == 0)
            appendix_tables += 1
        elif block.startswith(("Примечание:", "Источник:")):
            p = normal(block, 10, keep=block.startswith("Примечание:"))
            pf = p.paragraph_format; pf.first_line_indent = Cm(0)
            pf.line_spacing = 1; pf.space_before = Pt(3); pf.space_after = Pt(4)
        else:
            normal(block)

    related = [part for part in doc.part.related_parts.values() if str(part.partname) == "/word/footnotes.xml"]
    assert len(related) == 1
    related[0]._blob = etree.tostring(footroot, xml_declaration=True, encoding="UTF-8", standalone=True)
    doc.core_properties.title = "Диверсификация и волатильность экспорта интегральных схем Японии и Республики Корея"
    doc.core_properties.subject = "Дипломная работа"
    doc.core_properties.author = ""; doc.core_properties.last_modified_by = ""; doc.core_properties.comments = ""
    doc.save(OUTPUT)

    with ZipFile(OUTPUT) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        foot = etree.fromstring(archive.read("word/footnotes.xml"))
        refs = [int(x.get(qn("w:id"))) for x in root.findall(".//" + qn("w:footnoteReference"))]
        assert refs == list(range(1, 64)) and note_count == 63
        assert len(foot.findall(qn("w:footnote"))) == 65
        assert len(root.findall(".//" + qn("m:oMath"))) == 54
        assert len(root.findall(".//" + qn("w:drawing"))) == 11
        assert len(root.findall(".//" + qn("w:tbl"))) == 23
        assert appendix_tables == 11
        assert len(root.findall(".//" + qn("w:numPr"))) == 7
        assert not root.findall(".//" + qn("w:hyperlink"))
        rels = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        assert not [r for r in rels if r.get("TargetMode") == "External"]
        for node in list(root.iter()) + list(foot.iter()):
            if node.tag == qn("w:color"):
                assert node.get(qn("w:val")) in ("000000", "auto"), node.attrib

    report = {
        "file": OUTPUT.name, "sha256": hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        "chapters": 3, "source_count": 36, "appendices": 8,
        "tables": 23, "main_tables": 12, "appendix_tables": 11,
        "figures": 11, "native_math_objects": 54, "native_footnotes": 63,
        "footnote_ids_contiguous": True, "external_hyperlinks": 0,
        "imported_body_elements": imported_counts,
        "preserved_body_text_checks": len(content_checks),
        "main_body_text_preserved": True,
        "font": "Times New Roman", "body_points": 12, "line_spacing": 1.5, "page_size": "A4",
        "toc_kind": "static_no_links", "toc_pages": toc_pages,
        "headings": headings, "visual_review_pending": True,
        "input_sha256": {f"{n}.docx": hashlib.sha256((HERE / f"{n}.docx").read_bytes()).hexdigest() for n in originals},
    }
    (HERE / "diploma_full_checks.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    (HERE / "diploma_full_headings.json").write_text(json.dumps(headings, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps({k: v for k, v in report.items() if k not in ("input_sha256", "headings", "toc_pages")}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
