"""Build Chapter 1 with the typography of the supplied Naumov coursework.

Use the Codex primary runtime Python. The chapter has native Word footnotes
and editable equations, black typography, no tables, figures or hyperlinks.
"""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
import json
import re

from lxml import etree
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import CONTENT_TYPE as CT, RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


HERE = Path(__file__).resolve().parent
SOURCE_FILE = HERE.parent / 'literature' / 'sources.json'
OUTPUT = HERE / 'chapter1.docx'
BLACK = RGBColor(0, 0, 0)
NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
TOKEN = re.compile(r'\[\[ref ([^|]+)\|([^\]]+)\]\]')


def element(tag, **attrs):
    node = OxmlElement(tag)
    for name, value in attrs.items():
        node.set(qn(name), str(value))
    return node


def font(run_font, size=12, bold=False):
    run_font.name = 'Times New Roman'
    run_font.size = Pt(size)
    run_font.bold = bold
    run_font.color.rgb = BLACK


def add_text_node(parent, value):
    node = element('w:t')
    node.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    node.text = value
    parent.append(node)


def add_math_run(math, text):
    run = element('m:r')
    props = element('m:rPr')
    props.append(element('m:sty', **{'m:val': 'p'}))
    run.append(props)
    wt = element('m:t')
    wt.text = text
    run.append(wt)
    math.append(run)


def add_indexed(math, base, index, power=None):
    node = element('m:sSubSup' if power is not None else 'm:sSub')
    expr = element('m:e'); add_math_run(expr, base); node.append(expr)
    sub = element('m:sub'); add_math_run(sub, index); node.append(sub)
    if power is not None:
        sup = element('m:sup'); add_math_run(sup, power); node.append(sup)
    math.append(node)


def add_sum(math, index, body):
    node = element('m:nary')
    props = element('m:naryPr')
    props.append(element('m:chr', **{'m:val': '∑'}))
    props.append(element('m:limLoc', **{'m:val': 'subSup'}))
    props.append(element('m:supHide', **{'m:val': '1'}))
    node.append(props)
    sub = element('m:sub'); add_math_run(sub, index); node.append(sub)
    node.append(element('m:sup'))
    expr = element('m:e'); body(expr); node.append(expr)
    math.append(node)


def add_equation(doc, key):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mp = element('m:oMathPara')
    pr = element('m:oMathParaPr')
    pr.append(element('m:jc', **{'m:val': 'center'}))
    mp.append(pr)
    math = element('m:oMath')
    if key == 'variance':
        add_math_run(math, 'Var(G) = ')
        def own(e):
            add_indexed(e, 'w', 'i', '2'); add_math_run(e, ' Var(')
            add_indexed(e, 'g', 'i'); add_math_run(e, ')')
        add_sum(math, 'i', own)
        add_math_run(math, ' + 2 ')
        def pair(e):
            add_indexed(e, 'w', 'i'); add_math_run(e, ' '); add_indexed(e, 'w', 'j')
            add_math_run(e, ' Cov('); add_indexed(e, 'g', 'i')
            add_math_run(e, ', '); add_indexed(e, 'g', 'j'); add_math_run(e, ')')
        add_sum(math, 'i<j', pair)
    elif key == 'hhi':
        add_math_run(math, 'HHI = ')
        add_sum(math, 'j', lambda e: add_indexed(e, 's', 'j', '2'))
        add_math_run(math, ',   ')
        add_sum(math, 'j', lambda e: add_indexed(e, 's', 'j'))
        add_math_run(math, ' = 1')
    elif key == 'weights':
        add_indexed(math, 'w', 'k,j'); add_math_run(math, ' = ')
        add_indexed(math, 'a', 'k'); add_math_run(math, ' '); add_indexed(math, 'b', 'j|k')
        add_math_run(math, ',   ')
        add_sum(math, 'k', lambda e: add_indexed(e, 'a', 'k')); add_math_run(math, ' = 1,   ')
        add_sum(math, 'j', lambda e: add_indexed(e, 'b', 'j|k')); add_math_run(math, ' = 1')
    else:
        raise ValueError(key)
    mp.append(math)
    p._p.append(mp)


def main():
    rows = json.loads(SOURCE_FILE.read_text())['sources']
    sources = {r['id']: r for r in rows}
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin = Cm(2.5), Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.header_distance = sec.footer_distance = Cm(1)
    for style in doc.styles:
        if hasattr(style, 'font'):
            font(style.font)
        for n in list(style.element.iter(qn('w:pBdr'))):
            n.getparent().remove(n)
    normal = doc.styles['Normal']
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size in [('Title', 14), ('Heading 1', 13)]:
        style = doc.styles[name]
        font(style.font, size, True)
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(14 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(12)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.3
    doc.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name, kind in [('Footnote Text', WD_STYLE_TYPE.PARAGRAPH), ('Footnote Reference', WD_STYLE_TYPE.CHARACTER)]:
        if name not in doc.styles:
            doc.styles.add_style(name, kind)
    footstyle = doc.styles['Footnote Text']
    font(footstyle.font, 10)
    footstyle.paragraph_format.line_spacing = 1
    footstyle.paragraph_format.space_after = Pt(2)
    footstyle.paragraph_format.first_line_indent = Cm(0)
    footstyle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    font(doc.styles['Footnote Reference'].font, 9)
    doc.styles['Footnote Reference'].font.superscript = True
    footroot = etree.Element(qn('w:footnotes'), nsmap={'w': NS})
    for fid, type_name in [(-1, 'separator'), (0, 'continuationSeparator')]:
        foot = element('w:footnote', **{'w:type': type_name, 'w:id': fid})
        para = element('w:p'); run = element('w:r')
        run.append(element('w:' + type_name))
        para.append(run); foot.append(para); footroot.append(foot)
    used = set()
    footcount = 0

    def citation(p, source_id, locator):
        nonlocal footcount
        assert source_id in sources, source_id
        used.add(source_id); footcount += 1
        run = p.add_run()
        rp = run._r.get_or_add_rPr()
        rp.append(element('w:rStyle', **{'w:val': doc.styles['Footnote Reference'].style_id}))
        run._r.append(element('w:footnoteReference', **{'w:id': footcount}))
        foot = element('w:footnote', **{'w:id': footcount})
        para = element('w:p')
        props = element('w:pPr')
        props.append(element('w:pStyle', **{'w:val': footstyle.style_id}))
        para.append(props)
        r = element('w:r'); rprops = element('w:rPr')
        rprops.append(element('w:rStyle', **{'w:val': doc.styles['Footnote Reference'].style_id}))
        r.append(rprops); r.append(element('w:footnoteRef')); para.append(r)
        text_run = element('w:r')
        text = sources[source_id]['citation'].rstrip('.') + '. ' + locator.strip().rstrip('.') + '.'
        add_text_node(text_run, ' ' + text)
        para.append(text_run); foot.append(para); footroot.append(foot)

    for block in HERE.joinpath('chapter1_content.md').read_text().strip().split('\n\n'):
        block = block.strip()
        if block.startswith('# '):
            doc.add_paragraph(block[2:], style='Title')
        elif block.startswith('## '):
            doc.add_paragraph(block[3:], style='Heading 1')
        elif block.startswith('[[EQ '):
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            add_equation(doc, block[5:-2])
        else:
            p = doc.add_paragraph()
            pos = 0
            previous_ref = False
            for match in TOKEN.finditer(block):
                segment = block[pos:match.start()]
                if segment:
                    p.add_run(segment)
                elif previous_ref:
                    comma = p.add_run(',')
                    font(comma.font, 9)
                    comma.font.superscript = True
                citation(p, match.group(1), match.group(2))
                pos = match.end(); previous_ref = True
            p.add_run(block[pos:])

    bibliography_heading = doc.add_paragraph('Список использованных источников к первой главе', style='Heading 1')
    bibliography_heading.paragraph_format.page_break_before = True
    for i, sid in enumerate(sorted(used, key=lambda sid: sources[sid]['citation'].casefold()), 1):
        row = sources[sid]
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.left_indent = Cm(0.7)
        text = f"{i}. {row['citation']}"
        if row['doi']:
            text += f" DOI: {row['doi']}."
        p.add_run(text)

    fp = Part(PackURI('/word/footnotes.xml'), CT.WML_FOOTNOTES,
              etree.tostring(footroot, xml_declaration=True, encoding='UTF-8', standalone=True),
              doc.part.package)
    doc.part.relate_to(fp, RT.FOOTNOTES)
    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.first_line_indent = Cm(0)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = element('w:fldSimple', **{'w:instr': 'PAGE'})
    run = element('w:r'); add_text_node(run, '1'); field.append(run)
    footer._p.append(field)
    settings = doc.settings.element
    settings.append(element('w:updateFields', **{'w:val': 'true'}))
    lang = normal.element.get_or_add_rPr()
    lang.append(element('w:lang', **{'w:val': 'ru-RU'}))
    doc.core_properties.title = 'Глава 1 Теоретические основы диверсификации и волатильности экспорта интегральных схем'
    doc.core_properties.subject = 'Япония и Республика Корея'
    doc.core_properties.author = ''
    # Theme font attributes can override explicit names in LibreOffice and Word.
    for style in doc.styles:
        if not hasattr(style, 'font'):
            continue
        pr = style.element.get_or_add_rPr()
        fonts = pr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = element('w:rFonts'); pr.insert(0, fonts)
        for name in list(fonts.attrib):
            if 'theme' in etree.QName(name).localname.lower():
                del fonts.attrib[name]
        for name in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + name), 'Times New Roman')
        for child in list(pr):
            if child.tag in (qn('w:spacing'), qn('w:kern')):
                pr.remove(child)
        size = pr.find(qn('w:sz'))
        if size is not None:
            csize = pr.find(qn('w:szCs'))
            if csize is None:
                csize = element('w:szCs'); pr.append(csize)
            csize.set(qn('w:val'), size.get(qn('w:val')))
    doc.save(OUTPUT)
    with ZipFile(OUTPUT) as archive:
        xml = archive.read('word/document.xml')
        assert b'<w:tbl>' not in xml
        assert b'<w:drawing>' not in xml
        assert b'<w:hyperlink' not in xml
        assert xml.count(b'<w:footnoteReference ') == footcount
        assert b'[[ref ' not in xml and b'[[EQ ' not in xml
    report = {'document': OUTPUT.name, 'source_count': len(used), 'native_footnotes': footcount,
              'sections': 3, 'editable_equations': 3, 'tables': 0, 'figures': 0,
              'hyperlinks': 0, 'font': 'Times New Roman', 'body_points': 12,
              'body_line_spacing': 1.5, 'footnote_points': 10,
              'reference': 'Supplied Naumov coursework: formal prose, numbered footnotes, black typography'}
    (HERE/'chapter1_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report, ensure_ascii=False))


if __name__ == '__main__':
    main()
