"""Build the supervisor document using the inspected diploma3 typography.

Run with the Codex primary runtime Python. Tables read the frozen numerical
outputs; this command performs no new estimation or data collection.
"""
from __future__ import annotations
from pathlib import Path
import re
import json
from datetime import datetime, timezone
from zipfile import ZipFile

import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
RISK = HERE.parent/'risk_extension'
MONTHLY = RISK/'results/monthly'
POST = RISK/'post_estimation/results'
ASSETS = HERE/'topic2_assets'
BLACK = RGBColor(0, 0, 0)
COUNTRIES = {392:'Япония', 410:'Республика Корея'}


def num(x, digits=2, signed=False):
    if pd.isna(x):
        return 'Не определено'
    text = f'{x:+.{digits}f}' if signed else f'{x:.{digits}f}'
    return text.replace('-', '−').replace('.', ',')


def ci_text(point, low, high):
    return f'{num(point, signed=True)}\n[{num(low)}; {num(high)}]'


def xml(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    return node


def set_font(font, size=12, bold=None):
    font.name = 'Times New Roman'
    font.size = Pt(size)
    font.color.rgb = BLACK
    if bold is not None:
        font.bold = bold


def ptext(doc, text, style=None, size=12, bold=False, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.widow_control = True
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_font(r.font, size, bold)
    return p


def caption(doc, text):
    p = ptext(doc, text, size=11)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(4)
    return p


def table(doc, title, headers, rows, widths, note):
    caption(doc, title)
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    pr = t._tbl.tblPr
    pr.append(xml('w:tblW', **{'w:w':round(sum(widths)*567), 'w:type':'dxa'}))
    borders = xml('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(xml('w:'+side, **{'w:val':'single', 'w:sz':'4', 'w:color':'D9D9D9'}))
    pr.append(borders)
    margins = xml('w:tblCellMar')
    for side in ('top', 'bottom'):
        margins.append(xml('w:'+side, **{'w:w':80, 'w:type':'dxa'}))
    for side in ('left', 'right'):
        margins.append(xml('w:'+side, **{'w:w':95, 'w:type':'dxa'}))
    pr.append(margins)
    t.rows[0]._tr.get_or_add_trPr().append(xml('w:tblHeader'))
    for values in [headers, *rows]:
        cells = t.rows[0].cells if values is headers else t.add_row().cells
        for j, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell._tc.get_or_add_tcPr().append(xml('w:shd', **{'w:fill':'FFFFFF', 'w:val':'clear'}))
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.widow_control = True
            p.paragraph_format.keep_with_next = True
            r = p.add_run(str(value))
            set_font(r.font, 11, values is headers)
        cells[0]._tc.getparent().get_or_add_trPr().append(xml('w:cantSplit'))
    for j, column in enumerate(t.columns):
        column.width = Cm(widths[j])
    p = ptext(doc, note, size=10)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    return t


def mr(text):
    r = xml('m:r')
    rp = xml('w:rPr')
    rp.append(xml('w:rFonts', **{'w:ascii':'Cambria Math', 'w:hAnsi':'Cambria Math'}))
    rp.append(xml('w:sz', **{'w:val':24}))
    rp.append(xml('w:color', **{'w:val':'000000'}))
    r.append(rp)
    t = xml('m:t')
    t.text = text
    r.append(t)
    return r


def seq(*nodes):
    return [mr(n) if isinstance(n, str) else n for n in nodes]


def sub(base, lower):
    n = xml('m:sSub')
    e = xml('m:e'); e.append(mr(base)); n.append(e)
    s = xml('m:sub'); s.append(mr(lower)); n.append(s)
    return n


def sup(base, upper):
    n = xml('m:sSup')
    e = xml('m:e')
    e.append(mr(base) if isinstance(base, str) else base)
    n.append(e)
    s = xml('m:sup'); s.append(mr(upper)); n.append(s)
    return n


def fraction(top, bottom):
    n = xml('m:f')
    for name, nodes in [('m:num', top), ('m:den', bottom)]:
        part = xml(name)
        for node in seq(*nodes): part.append(node)
        n.append(part)
    return n


def equation(doc, name):
    if name == 'growth':
        nodes = seq(sub('g','i,t'), ' = ', fraction(
            ['2(', sub('x','i,t'), ' − ', sub('x','i,t−12'), ')'],
            [sub('x','i,t'), ' + ', sub('x','i,t−12')]))
    elif name == 'weight':
        nodes = seq(sub('a','g,d'), ' = ', sub('w','g'), sub('p','d|g'))
    elif name == 'hhi':
        nary = xml('m:nary')
        pr = xml('m:naryPr')
        pr.append(xml('m:chr', **{'m:val':'∑'}))
        pr.append(xml('m:limLoc', **{'m:val':'subSup'}))
        pr.append(xml('m:supHide', **{'m:val':'1'}))
        nary.append(pr)
        lower=xml('m:sub'); lower.append(mr('d')); nary.append(lower)
        nary.append(xml('m:sup'))
        body=xml('m:e'); body.append(sup(sub('s','d'),'2')); nary.append(body)
        nodes=seq('HHI = ', nary)
    elif name == 'variance':
        radical=xml('m:rad')
        pr=xml('m:radPr');pr.append(xml('m:degHide', **{'m:val':'1'}));radical.append(pr)
        radical.append(xml('m:deg'))
        e=xml('m:e');e.append(mr('V'));radical.append(e)
        nodes=seq('V = ',sup('a','T'),'Σa,     S = 100',radical)
    else:
        nodes=seq(sub('φ','w' if name == 'shapley_product' else 'p'),' = ',fraction(['1'],['2']),' [')
        pairs = [('10','00'),('11','01')] if name == 'shapley_product' else [('01','00'),('11','10')]
        for j, (a,b) in enumerate(pairs):
            if j:nodes.extend(seq(' + '))
            nodes.extend(seq('(',sub('V',a),' − ',sub('V',b),')'))
        nodes.extend(seq(']'))
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(3)
    p.paragraph_format.space_after=Pt(8)
    om=xml('m:oMath');
    for node in nodes:om.append(node)
    omp=xml('m:oMathPara');omp.append(om);p._p.append(omp)
    return p


def make_figures():
    ASSETS.mkdir(exist_ok=True)
    plt.rcParams.update({'font.family':'DejaVu Serif', 'font.size':10,
        'axes.spines.top':False, 'axes.spines.right':False})
    observed=pd.read_csv(MONTHLY/'observed_growth.csv')
    fig,ax=plt.subplots(figsize=(6.4,2.65),layout='constrained')
    for reporter,style in [(392,'-'),(410,'--')]:
        d=observed[observed.reporter.eq(reporter)&~observed.include_parts&observed.lag.eq(12)]
        ax.plot(pd.to_datetime(d.period.astype(str),format='%Y%m'),d.growth_pp,style,
                color='black',linewidth=1.05,label=COUNTRIES[reporter])
    ax.axhline(0,color='.55',linewidth=.6)
    ax.set(ylabel='Симметричный рост г/г, %',xlabel='Месяц')
    ax.grid(axis='y',color='.88',linewidth=.5)
    ax.legend(frameon=False,ncol=2,fontsize=10,loc='lower left')
    fig.savefig(ASSETS/'observed.png',dpi=240,facecolor='white')
    plt.close(fig)
    loo=pd.read_csv(POST/'leave_one_year_out.csv')
    summary=pd.read_csv(MONTHLY/'risk_summary.csv')
    summary=summary[summary.case.eq('primary')].set_index('reporter')
    fig,axes=plt.subplots(1,2,figsize=(6.4,2.6),layout='constrained')
    for ax,(reporter,name) in zip(axes,COUNTRIES.items()):
        d=loo[loo.case.eq('primary')&loo.reporter.eq(reporter)]
        ax.plot(d.omitted_year,d.delta_sd_pp,'-o',color='black',markersize=3,linewidth=.9)
        ax.axhline(summary.loc[reporter,'delta_sd_pp'],color='.35',linestyle='--',linewidth=.8)
        ax.axhline(0,color='.6',linewidth=.6)
        ax.set(title=name,xlabel='Исключенный год',xticks=[2018,2020,2022,2024])
        ax.grid(axis='y',color='.88',linewidth=.5)
    axes[0].set_ylabel('Изменение SD, п.п.')
    fig.savefig(ASSETS/'omission.png',dpi=240,facecolor='white')
    plt.close(fig)


def add_figure(doc, name):
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_after=Pt(3)
    shape=p.add_run().add_picture(str(ASSETS/f'{name}.png'),width=Cm(16.5))
    shape._inline.docPr.set('descr', 'График фактического роста экспорта двух стран' if name=='observed'
                            else 'Изменение условного стандартного отклонения при исключении каждого года')
    title=('Рисунок 1 Фактический рост стоимости экспорта готовых интегральных схем'
           if name=='observed' else 'Рисунок 2 Чувствительность к исключению одного года')
    cap=ptext(doc,title,size=11)
    cap.paragraph_format.keep_with_next=True
    note=('Источник: UN Comtrade, расчеты автора. Четыре HS6, январь 2018 — декабрь 2024 года.'
          if name=='observed' else 'Источник: расчеты автора. Пунктир — полная основная оценка. Доли 2017/2024 и сетка фиксированы. Дополнительная диагностика после основной оценки.')
    p=ptext(doc,note,size=10)
    p.paragraph_format.line_spacing=1.0


def add_table(doc, key, data):
    summary,states,intervals,annual,memory,covariance,destinations=data
    primary=summary[summary.case.eq('primary')].set_index('reporter')
    if key=='goods':
        table(doc,'Таблица 1 Товарные категории исследования',
            ['Код HS','Категория','Использование'],[
                ['854231','Процессоры и контроллеры','Основная оценка'],
                ['854232','Память','Основная оценка'],
                ['854233','Усилители','Основная оценка'],
                ['854239','Прочие интегральные схемы','Основная оценка'],
                ['854290','Части','Отдельная проверка'],
                ['8542','Интегральные схемы и части','Контроль суммы пяти HS6']],
            [2.4,7.0,7.1],'Источник: UN Comtrade и номенклатура HS. Основной риск оценивается для четырех категорий готовых схем.')
    elif key=='endpoints':
        rows=[]
        for reporter in COUNTRIES:
            for year in [2017,2024]:
                r=annual[annual.reporter.eq(reporter)&annual.year.eq(year)&~annual.include_parts].iloc[0]
                rows.append([COUNTRIES[reporter],year,num(r.export_usd/1e9),num(100*r.memory_share),num(r.hhi,6)])
        table(doc,'Таблица 2 Стоимость экспорта и показатели его товарной и географической структуры',
            ['Экспортер','Год','Экспорт, млрд долл.','Память, %','HHI'],rows,
            [3.5,1.2,3.9,3.6,4.3],
            'Источник: UN Comtrade, расчеты автора по месячным данным. Четыре категории готовых схем; HHI по всем направлениям, шкала 0–1.')
    elif key=='states':
        names={'initial':'Доли и география 2017', 'product_change':'Только товарные доли 2024',
               'geography_change':'Только география 2024', 'final':'Доли и география 2024'}
        rows=[]
        for state,label in names.items():
            vals=[]
            for reporter in COUNTRIES:
                r=states[states.case.eq('primary')&states.reporter.eq(reporter)&states.state.eq(state)].iloc[0]
                vals.append(num(r.sd_pp))
            rows.append([label,*vals])
        table(doc,'Таблица 3 Условное стандартное отклонение четырех структур',
            ['Состояние','Япония, п.п.','Корея, п.п.'],rows,[8.1,4.2,4.2],
            'Источник: расчеты автора. Для каждого экспортера одна ковариация 2018–2024 годов; 84 месячных годовых темпа в каждом состоянии.')
    elif key=='primary_effects':
        rows=[]
        for reporter in COUNTRIES:
            r=primary.loc[reporter];vals=[]
            for metric in ['delta_sd_pp','mix_variance_pp2','geography_variance_pp2']:
                c=intervals[intervals.reporter.eq(reporter)&intervals.metric.eq(metric)&intervals.block_months.eq(12)].iloc[0]
                vals.append(ci_text(r[metric],c.ci_low,c.ci_high))
            rows.append([COUNTRIES[reporter],*vals])
        table(doc,'Таблица 4 Основное изменение и вклады факторов',
            ['Экспортер','Изменение SD, п.п.','Товарный вклад, п.п.²','География, п.п.²'],rows,
            [3.1,4.0,4.7,4.7],
            'Источник: расчеты автора. В квадратных скобках 95%-интервалы; 1 999 парных блочных повторов, блок 12 месяцев. Вклады относятся к дисперсии.')
    elif key=='memory':
        rows=[]
        for factor,label in [('memory_share','Доля памяти'),('within_nonmemory_mix','Состав остальных категорий'),('geography','География')]:
            vals=[]
            for reporter in COUNTRIES:
                r=memory[memory.case.eq('primary')&memory.reporter.eq(reporter)&memory.factor.eq(factor)&memory.block_months.eq(12)].iloc[0]
                vals.append(ci_text(r.estimate_variance_pp2,r.ci_low,r.ci_high))
            rows.append([label,*vals])
        table(doc,'Таблица 5 Дополнительное разложение товарного вклада',
            ['Часть изменения','Япония, п.п.²','Корея, п.п.²'],rows,[6.1,5.2,5.2],
            'Источник: расчеты автора. Поисковая диагностика после основной оценки. 95%-интервалы при блоке 12 месяцев; без поправки на множество сравнений.')
    elif key=='covariance':
        rows=[]
        for component,label in [('own_variances','Собственные дисперсии'),('cross_covariances','Перекрестные ковариации')]:
            r=covariance[covariance.case.eq('primary')&covariance.reporter.eq(410)&covariance.factor.eq('geography')&covariance.component.eq(component)&covariance.block_months.eq(12)].iloc[0]
            rows.append([label,num(r.estimate_variance_pp2,signed=True),f'[{num(r.ci_low)}; {num(r.ci_high)}]'])
        r=intervals[intervals.reporter.eq(410)&intervals.metric.eq('geography_variance_pp2')&intervals.block_months.eq(12)].iloc[0]
        rows.append(['Общий географический вклад',num(primary.loc[410,'geography_variance_pp2'],signed=True),f'[{num(r.ci_low)}; {num(r.ci_high)}]'])
        table(doc,'Таблица 6 Компоненты географического вклада Кореи',
            ['Компонент','Вклад, п.п.²','95%-интервал'],rows,[7.1,3.6,5.8],
            'Источник: расчеты автора. Дополнительная диагностика. Точечные компоненты складываются в общий вклад; границы интервалов не складываются.')
    elif key=='destinations':
        names={156:'Китай',344:'Гонконг',490:'Other Asia, nes',704:'Вьетнам',702:'Сингапур',608:'Филиппины',842:'США',764:'Таиланд',-1:'Остальные направления'}
        rows=[]
        for code,label in names.items():
            r=destinations[destinations.case.eq('primary')&destinations.reporter.eq(410)&destinations.factor.eq('geography')&destinations.destination.eq(code)&destinations.block_months.eq(12)].iloc[0]
            rows.append([label,num(r.estimate_variance_pp2,signed=True),f'[{num(r.ci_low)}; {num(r.ci_high)}]'])
        table(doc,'Таблица 7 Географический вклад Кореи по направлениям',
            ['Направление','Вклад, п.п.²','95%-интервал'],rows,[6.1,4.0,6.4],
            'Источник: расчеты автора. Дополнительная симметричная атрибуция с учетом ковариаций, блок 12 месяцев. В сумме точечные значения дают +98,84 п.п.².')
    elif key=='robustness':
        labels={'primary':'Основная оценка','top5':'Пять кандидатов направлений','top12':'Двенадцать кандидатов направлений',
                'include_parts':'С частями HS854290','cn_hk_combined':'Китай и Гонконг вместе',
                'monthly_lag1':'Симметричный месячный рост', 'quarterly_yoy':'Симметричный квартальный рост г/г',
                'ordinary_pct_yoy':'Обычный процентный рост г/г','sigma_2018_2020':'Ковариация 2018–2020','sigma_2021_2024':'Ковариация 2021–2024'}
        rows=[]
        for case,label in labels.items():
            vals=[]
            for reporter in COUNTRIES:
                r=summary[summary.case.eq(case)&summary.reporter.eq(reporter)].iloc[0]
                vals.append(num(r.delta_sd_pp,signed=True) if r.status=='ok' else 'Не определено')
            rows.append([label,*vals])
        table(doc,'Таблица 8 Изменение условного стандартного отклонения в разных вариантах',
            ['Вариант','Япония, п.п.','Корея, п.п.'],rows,[9.1,3.7,3.7],
            'Источник: расчеты автора. Во всех строках доли 2017/2024; окна и определения роста различаются. Неопределенность связана с нулевыми знаменателями, а не с неполной панелью.')
    else:
        raise KeyError(key)


def main():
    quality=json.loads((MONTHLY/'data_quality.json').read_text())
    validation=json.loads((MONTHLY/'analysis_validation.json').read_text())
    assert quality['analysis_ready'] and quality['monthly_panel_sha256']==validation['monthly_panel_sha256']
    make_figures()
    data=[pd.read_csv(path) for path in [MONTHLY/'risk_summary.csv', MONTHLY/'risk_states.csv',
          MONTHLY/'risk_intervals.csv',MONTHLY/'annual_indicators_from_months.csv',
          POST/'memory_share_decomposition.csv',POST/'covariance_attribution_intervals.csv',
          POST/'destination_attribution_intervals.csv']]
    doc=Document()
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.left_margin=Cm(2.5);section.right_margin=Cm(2)
    section.top_margin=Cm(2);section.bottom_margin=Cm(2)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
        if hasattr(style,'font'):
            set_font(style.font)
            rp=style.element.get_or_add_rPr()
            rf=rp.find(qn('w:rFonts'))
            if rf is None:rf=xml('w:rFonts');rp.append(rf)
            for attr in list(rf.attrib):
                if 'Theme' in attr or 'theme' in attr:rf.attrib.pop(attr)
            for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')
    normal=doc.styles['Normal']
    normal.paragraph_format.line_spacing=1.15
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.first_line_indent=Cm(0)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    for name in ['Heading 1','Heading 2']:
        style=doc.styles[name];set_font(style.font,12,True)
        style.paragraph_format.space_before=Pt(6)
        style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.keep_with_next=True
    title=doc.styles['Title'];set_font(title.font,12,True)
    title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing=1.0
    title.paragraph_format.space_after=Pt(9)
    source=(HERE/'topic2_content.md').read_text()
    for block in re.split(r'\n\s*\n',source.strip()):
        if block.startswith('# '):
            p=ptext(doc,block[2:],style='Title',bold=True,align=WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=True
        elif block.startswith('## '):
            ptext(doc,block[3:],style='Heading 1',bold=True)
        elif block.startswith('[[TABLE '):add_table(doc,block[8:-2],data)
        elif block.startswith('[[EQ '):equation(doc,block[5:-2])
        elif block.startswith('[[FIG '):add_figure(doc,block[6:-2])
        else:ptext(doc,block.replace('\n',' '))
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();set_font(run.font,11)
    run._r.append(xml('w:fldChar', **{'w:fldCharType':'begin'}))
    instruction=xml('w:instrText');instruction.text=' PAGE ';run._r.append(instruction)
    run._r.append(xml('w:fldChar', **{'w:fldCharType':'separate'}))
    text=xml('w:t');text.text='1';run._r.append(text)
    run._r.append(xml('w:fldChar', **{'w:fldCharType':'end'}))
    doc.core_properties.title='Волатильность экспорта интегральных схем Японии и Республики Корея с учетом товарной структуры и географии поставок'
    doc.core_properties.subject='Данные и результаты исследования для научной руководительницы'
    doc.core_properties.author='Артём';doc.core_properties.last_modified_by='Артём'
    doc.core_properties.created=datetime(2026,10,16,tzinfo=timezone.utc)
    doc.core_properties.modified=datetime.now(timezone.utc)
    path=HERE/'topic2.docx';doc.save(path)
    with ZipFile(path) as z:
        body=z.read('word/document.xml').decode()
        assert '<w:hyperlink' not in body and '[[TABLE' not in body and '[[EQ' not in body and '[[FIG' not in body
        assert body.count('<m:oMath>')==6
        assert '<w:pBdr' not in body and '<w:pBdr' not in z.read('word/styles.xml').decode()
        assert not any('hyperlink' in z.read(n).decode(errors='ignore') for n in z.namelist() if n.endswith('.rels'))
    print(json.dumps(dict(path=str(path),source_words=len(source.split()),tables=len(doc.tables),
        figures=len(doc.inline_shapes),editable_equations=6,hyperlinks=0,
        font='Times New Roman 12 pt',line_spacing=1.15,paper='A4',margins_cm=[2.5,2,2,2]),ensure_ascii=False,indent=2))


if __name__=='__main__':main()
