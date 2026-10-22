"""Replace four reviewed images; split Б.1 with a numbered continuation."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from io import BytesIO
import argparse,json,hashlib,re
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from PIL import Image

HERE=Path(__file__).resolve().parent
def sha(data):return hashlib.sha256(data).hexdigest()
def main():
 ap=argparse.ArgumentParser();ap.add_argument('--base',type=Path,required=True);ap.add_argument('--manifest',type=Path,required=True);ap.add_argument('--assets',type=Path,required=True);ap.add_argument('--output',type=Path,default=HERE/'diploma_full.docx');a=ap.parse_args()
 report=json.loads(a.manifest.read_text());assert sha(a.base.read_bytes())==report['output_sha256'];d=Document(a.base)
 assert len(d.inline_shapes)==32;old_text=[p.text for p in d.paragraphs]
 files={10:['10_product_destination.png'],21:['21_cell_expansion.png'],24:['24_destination_correlations_part1.png','24_destination_correlations_part2.png'],23:['23_monthly_product_growth.png']}
 def dimensions(path):
  w,h=Image.open(path).size;return Cm(16.5),Cm(16.5*h/w)
 def attach_dimensions(inline,path):
  w,h=dimensions(path)
  for node in inline.xpath('.//wp:extent | .//a:xfrm/a:ext'):
   node.set('cx',str(int(w)));node.set('cy',str(int(h)))
 def after(anchor,template,text=None,picture=None):
  node=OxmlElement('w:p')
  if template._p.pPr is not None:node.append(deepcopy(template._p.pPr))
  anchor._p.addnext(node);p=Paragraph(node,anchor._parent)
  if picture:
   w,h=dimensions(picture);shape=p.add_run().add_picture(str(picture),width=w,height=h);shape._inline.docPr.set('descr','Рисунок Б.1. Корреляции направлений: Республика Корея (продолжение)')
  else:
   run=p.add_run(text)
   if template.runs and template.runs[0]._r.rPr is not None:run._r.insert(0,deepcopy(template.runs[0]._r.rPr))
  return p
 original_shapes=list(d.inline_shapes);flattened=[(f,item) for f in report['figures'] for item in f['images']]
 for shape,(figure,item) in zip(original_shapes,flattened):
  n=figure['archive_number']
  if n not in files:continue
  part=d.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.embed]
  assert sha(part.blob)==item['sha256']
  path=a.assets/files[n][0];part._blob=path.read_bytes();attach_dimensions(shape._inline,path)
  if n==24:
   ps=list(d.paragraphs);ci=next(i for i,p in enumerate(ps) if p.text.startswith('Рисунок Б.1 — '));imagep=ps[ci-1];caption=ps[ci];source=ps[ci+1]
   assert imagep._p.xpath('.//w:drawing') and source.text.startswith('Источник:')
   p=after(source,imagep,picture=a.assets/files[n][1]);p=after(p,caption,text='Продолжение рисунка Б.1');after(p,source,text=source.text)
  figure['images']=[{'file':name,'sha256':sha((a.assets/name).read_bytes()),'width_cm':16.5,'print_redesign':True} for name in files[n]]
 a.output.parent.mkdir(parents=True,exist_ok=True);d.save(a.output)
 with ZipFile(a.base) as b,ZipFile(a.output) as z:
  assert z.read('word/footnotes.xml')==b.read('word/footnotes.xml')
  assert z.read('word/document.xml').count(b'<m:oMath>')==54
 final=Document(a.output);assert len(final.inline_shapes)==33 and len(final.tables)==23
 remaining=[p.text for p in final.paragraphs];pos=0
 for t in old_text:
  pos=remaining.index(t,pos)+1
 report['output_sha256']=sha(a.output.read_bytes());report['image_panels']=33
 report['print_redesign']={'figures':['2.7','3.4','Б.1','Д.1'],'body_font_pt_min':8.5,'numeric_source_values':954,'data_changed':False,'correlation_separate_country_panels':True,'logical_figures':30}
 report.pop('rendered_pages',None);report.pop('toc_pages',None)
 (HERE/'figure_insertion_manifest.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
 (HERE/'four_figures_validation.json').write_bytes((a.assets/'four_figures_validation.json').read_bytes())
 print(json.dumps({'logical_figures':30,'image_panels':33,'original_text_preserved':True,'native_equations':54,'tables':23,'output_sha256':report['output_sha256']}))

if __name__=='__main__':main()
