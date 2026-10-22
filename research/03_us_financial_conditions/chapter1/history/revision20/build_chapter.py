from pathlib import Path
import re,json,subprocess,os,zipfile
from lxml import etree
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path('/workspace/scratch/3c2847f1995e/chapter1_work')
OUT=ROOT/'output';OUT.mkdir(exist_ok=True)
REF=json.loads((ROOT/'references.json').read_text())
text=(ROOT/'chapter1_draft.md').read_text()
groups=re.findall(r'\{\{(S\d+(?:;S\d+)*)\}\}',text)
used=[]
for group in groups:
 for sid in group.split(';'):
  if sid not in used:used.append(sid)
numbers={sid:i+1 for i,sid in enumerate(used)}

doc=Document()
sec=doc.sections[0]
sec.page_width,sec.page_height=Cm(21),Cm(29.7)
sec.left_margin,sec.right_margin=Cm(2.5),Cm(1.5)
sec.top_margin,sec.bottom_margin=Cm(2),Cm(2)
sec.footer_distance=Cm(1)
for name in ['Normal','Title','Heading 1','Heading 2','Footnote Text','Footnote Reference']:
 if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.CHARACTER if name=='Footnote Reference' else WD_STYLE_TYPE.PARAGRAPH)
 s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
 s.font.size=Pt(12)
 if s.type==WD_STYLE_TYPE.PARAGRAPH:
  s.paragraph_format.space_after=Pt(0);s.paragraph_format.space_before=Pt(0)
  s.paragraph_format.line_spacing=1.5;s.paragraph_format.widow_control=True
  s.paragraph_format.first_line_indent=Cm(1.25)
normal=doc.styles['Normal'];normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Title','Heading 1','Heading 2']:
 s=doc.styles[name];s.font.bold=True;s.font.size=Pt(13 if name=='Heading 2' else 14);s.paragraph_format.first_line_indent=Cm(0)
 s.paragraph_format.keep_with_next=True
 s.paragraph_format.space_before=Pt(12 if name=='Heading 2' else 0)
 s.paragraph_format.space_after=Pt(6)
 s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.styles['Heading 1'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Footnote Text'].font.size=Pt(10)
doc.styles['Footnote Text'].paragraph_format.first_line_indent=Cm(0)
doc.styles['Footnote Text'].paragraph_format.line_spacing=1
doc.styles['Footnote Reference'].font.size=Pt(10)
doc.styles['Footnote Reference'].font.superscript=True

for sty in doc.styles:
 if sty.type in (WD_STYLE_TYPE.PARAGRAPH,WD_STYLE_TYPE.CHARACTER):
  rf=sty._element.get_or_add_rPr().get_or_add_rFonts()
  for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme','csTheme']:rf.attrib.pop(qn('w:'+k),None)
  for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Times New Roman')
for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)

doc.core_properties.title='Глава 1 Теоретические основы международной передачи финансовых условий и прогнозирования рисков роста'
doc.core_properties.subject='Финансовые условия США в прогнозировании рисков роста еврозоны'
doc.core_properties.author=''
doc.core_properties.keywords='Growth at Risk; NFCI; ANFCI; еврозона; международные финансовые связи'

notes=[]
def footnote_repl(m):
 ids=m.group(1).split(';');notes.append(' '.join(REF[s]['chapter1_reference'] for s in ids))
 return f'[[FN{len(notes):03d}]]'

for block in re.split(r'\n\s*\n',text.strip()):
 block=block.strip()
 if block.startswith('# '):
  p=doc.add_paragraph(block[2:],style='Heading 1')
 elif block.startswith('## '):
  p=doc.add_paragraph(block[3:],style='Heading 2')
 elif block.startswith('$$'):
  m=re.fullmatch(r'\$\$(.*?)\$\$\s*(\(\d+\.\d+\))',block)
  if not m:raise ValueError(block)
  p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
  om=OxmlElement('m:oMath');mr=OxmlElement('m:r');prop=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');prop.append(sty);mr.append(prop)
  wr=OxmlElement('w:rPr');fs=OxmlElement('w:sz');fs.set(qn('w:val'),'24');wr.append(fs);mr.append(wr)
  mt=OxmlElement('m:t');mt.text=m.group(1)+'    '+m.group(2);mr.append(mt);om.append(mr);p._p.append(om)
 else:
  doc.add_paragraph(re.sub(r'\{\{(S\d+(?:;S\d+)*)\}\}',footnote_repl,block))

doc.add_page_break()
p=doc.add_paragraph('Список использованных источников к главе 1',style='Heading 1')
for sid in used:
 x=REF[sid]
 ref=x['chapter1_reference']
 if not x.get('doi') or sid in ['S02','S12','S14','S23','S24','S26','S30','S33','S34','S36','S37','S38','S41','S44']:
  ref+=' URL: '+x.get('url','')+' (дата обращения 16.10.2026).'
 p=doc.add_paragraph(f'{numbers[sid]}. '+ref)
 p.paragraph_format.first_line_indent=Cm(-0.75);p.paragraph_format.left_indent=Cm(0.75)
 p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3)
 p.paragraph_format.keep_together=True
 for run in p.runs:run.font.size=Pt(12)

foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
foot.paragraph_format.first_line_indent=Cm(0)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)

draft=ROOT/'chapter_with_markers.docx';doc.save(draft)
helper='/root/.codex/skills/builtins/documents/scripts/insert_note.py'
for i,note in enumerate(notes,1):
 tmp=ROOT/'chapter_note_tmp.docx'
 subprocess.run([os.environ['CODEX_PRIMARY_RUNTIME_PYTHON'],helper,str(draft),'--kind','footnote','--marker',f'[[FN{i:03d}]]','--text',note,'--out',str(tmp)],check=True,capture_output=True,text=True)
 os.replace(tmp,draft)

# Set explicit note typography in addition to the Word footnote styles.
final=OUT/'Chapter_1_US_Financial_Conditions_Euro_Area.docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(draft) as zi:
 data={i.filename:zi.read(i.filename) for i in zi.infolist()}
tree=etree.fromstring(data['word/footnotes.xml'])
for note in tree.findall('w:footnote',ns):
 note_id=int(note.get(qn('w:id')))
 if note_id<1:
  note.set(qn('w:type'),'separator' if note_id==-1 else 'continuationSeparator')
  continue
 for p in note.findall('w:p',ns):
  pp=p.find('w:pPr',ns)
  if pp is None:pp=etree.Element(qn('w:pPr'));p.insert(0,pp)
  st=pp.find('w:pStyle',ns)
  if st is None:st=etree.SubElement(pp,qn('w:pStyle'))
  st.set(qn('w:val'),'FootnoteText')
  keep=pp.find('w:keepLines',ns)
  if keep is None:etree.SubElement(pp,qn('w:keepLines'))
  sp=pp.find('w:spacing',ns)
  if sp is None:sp=etree.SubElement(pp,qn('w:spacing'))
  sp.set(qn('w:line'),'240');sp.set(qn('w:lineRule'),'auto');sp.set(qn('w:before'),'0');sp.set(qn('w:after'),'0')
  for r in p.findall('w:r',ns):
   rp=r.find('w:rPr',ns)
   if rp is None:rp=etree.Element(qn('w:rPr'));r.insert(0,rp)
   sz=rp.find('w:sz',ns)
   if sz is None:sz=etree.SubElement(rp,qn('w:sz'))
   sz.set(qn('w:val'),'20')
   rf=rp.find('w:rFonts',ns)
   if rf is None:rf=etree.SubElement(rp,qn('w:rFonts'))
   for k in ['ascii','hAnsi','cs','eastAsia']:rf.set(qn('w:'+k),'Times New Roman')
data['word/footnotes.xml']=etree.tostring(tree,encoding='UTF-8',xml_declaration=True,standalone=True)
with zipfile.ZipFile(final,'w',zipfile.ZIP_DEFLATED) as zo:
 for name,b in data.items():zo.writestr(name,b)

md=re.sub(r'\{\{(S\d+(?:;S\d+)*)\}\}',lambda m:'['+', '.join(str(numbers[z]) for z in m.group(1).split(';'))+']',text)
md+='\n\n## Список использованных источников к главе 1\n\n'
for sid in used:
 x=REF[sid];md+=f'{numbers[sid]}. {x["chapter1_reference"]} '+(f'https://doi.org/{x["doi"]}' if x.get('doi') else x.get('url',''))+'\n\n'
(OUT/'Chapter_1_US_Financial_Conditions_Euro_Area.md').write_text(md)
(ROOT/'chapter_metrics.json').write_text(json.dumps({'words_body':len(text.split()),'characters_body':len(text),'footnotes':len(notes),'references':len(used),'source_order':used},ensure_ascii=False,indent=2))
print(json.dumps({'docx':str(final),'body_words':len(text.split()),'footnotes':len(notes),'sources':len(used)},ensure_ascii=False))
