"""Plain student memo; editable text and ordinary Word tables."""
from pathlib import Path
import re,shutil
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

P=Path(__file__).resolve().parent
template=P/'Note_for_supervisor_before_revision16.docx'
doc=Document(template) if template.exists() else Document()
body=doc._element.body
for child in list(body):
 if child.tag!=qn('w:sectPr'):body.remove(child)
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.left_margin=Cm(2.5);sec.right_margin=Cm(2);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(12);style.font.color.rgb=RGBColor(0,0,0)
 style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
 style.paragraph_format.space_before=Pt(0);style.paragraph_format.space_after=Pt(5);style.paragraph_format.line_spacing=1.15
 ppr=style._element.find(qn('w:pPr'))
 if ppr is not None:
  for border in list(ppr.findall(qn('w:pBdr'))):ppr.remove(border)
doc.styles['Normal'].paragraph_format.first_line_indent=Cm(1.25)
doc.styles['Title'].font.size=Pt(14);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.first_line_indent=Cm(0);doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Title'].paragraph_format.keep_with_next=True
for name in ['Heading 1','Heading 2']:
 style=doc.styles[name];style.font.bold=True;style.paragraph_format.first_line_indent=Cm(0)
 style.paragraph_format.space_before=Pt(10);style.paragraph_format.space_after=Pt(5);style.paragraph_format.keep_with_next=True

def font(run,size=None):
 run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0)
 rpr=run._element.get_or_add_rPr();rfont=rpr.get_or_add_rFonts()
 for k in ['ascii','hAnsi','eastAsia','cs']:rfont.set(qn('w:'+k),'Times New Roman')
 if size:run.font.size=Pt(size)

def text(p,s,size=None):
 for frag in re.split(r'(https?://[^\s;]+)',s):
  if frag.startswith('http'):
   link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(frag,RT.HYPERLINK,is_external=True))
   r=OxmlElement('w:r');rp=OxmlElement('w:rPr');rf=OxmlElement('w:rFonts')
   for k in ['ascii','hAnsi','eastAsia']:rf.set(qn('w:'+k),'Times New Roman')
   rp.append(rf);c=OxmlElement('w:color');c.set(qn('w:val'),'000000');rp.append(c)
   z=OxmlElement('w:sz');z.set(qn('w:val'),str(int((size or 12)*2)));rp.append(z);r.append(rp)
   tx=OxmlElement('w:t');tx.text=frag;r.append(tx);link.append(r);p._p.append(link)
  else:font(p.add_run(frag),size)

def table(lines):
 rows=[[c.strip() for c in s.strip().strip('|').split('|')] for s in lines]
 rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c) for c in r)]
 n=len(rows[0]);t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[3.1,7.0,6.4] if n==3 and rows[0][0]=='Переменная' else [9.7,3.4,3.4] if n==3 else [5.7,2.7,2.7,2.7,2.7]
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 props=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'B7B7B7');borders.append(e)
 props.append(borders)
 for i,row in enumerate(rows):
  tr=t.add_row()
  pr=tr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for j,(cell,s) in enumerate(zip(tr.cells,row)):
   cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   cp=cell._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
   for edge,val in [('top',80),('bottom',80),('left',95),('right',95)]:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');m.append(e)
   cp.append(m);p=cell.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1.05
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j>0 and (n==5 or rows[0][0]=='Проверка h=1') else WD_ALIGN_PARAGRAPH.LEFT
   text(p,s,11)
   if i==0:
    for r in p.runs:r.bold=True
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY;p.paragraph_format.line_spacing=Pt(3);p.paragraph_format.first_line_indent=Cm(0);font(p.add_run(),2)

lines=(P/'RESPONSE_TO_SUPERVISOR.md').read_text().splitlines();i=0;front=True
while i<len(lines):
 s=lines[i].strip();i+=1
 if not s:continue
 if s.startswith('|'):
  group=[s]
  while i<len(lines) and lines[i].lstrip().startswith('|'):group.append(lines[i]);i+=1
  table(group);continue
 if s.startswith('# '):
  p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;text(p,s[2:],14);continue
 if s.startswith('## '):
  front=False;p=doc.add_paragraph(style='Heading 1');text(p,s[3:]);continue
 p=doc.add_paragraph();p.paragraph_format.widow_control=True
 if front:
  p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True;text(p,s)
 elif re.match(r'^\[(?:\d+|Д\d+)\]',s):
  p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(3.5);p.alignment=WD_ALIGN_PARAGRAPH.LEFT;text(p,s,10.5)
 else:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;text(p,s)

for section in doc.sections:
 for h in [section.header,section.footer]:
  for old in list(h._element):h._element.remove(old)
 footer=section.footer.add_paragraph();footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;footer.paragraph_format.first_line_indent=Cm(0)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.author='Рустам Чулиев';doc.core_properties.title='Финансовые условия США в прогнозировании рисков роста еврозоны'
doc.core_properties.subject='Уточнение постановки и результаты проверки замечаний научного руководителя'
doc.core_properties.comments=''
out=P/'Note_for_supervisor.docx';doc.save(out);print(out)
