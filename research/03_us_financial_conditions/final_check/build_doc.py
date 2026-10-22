from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
P=Path(__file__).resolve().parent;d=Document();s=d.sections[0];s.page_height=Cm(29.7);s.page_width=Cm(21);s.top_margin=s.bottom_margin=Cm(2);s.left_margin=Cm(2.5);s.right_margin=Cm(2)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(12);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.1
 st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
d.styles['Title'].font.size=Pt(16);d.styles['Title'].font.bold=True;d.styles['Heading 1'].font.size=Pt(13);d.styles['Heading 1'].font.bold=True
for line in (P/'DECISION.md').read_text().splitlines():
 if not line:continue
 style='Title' if line.startswith('# ') else 'Heading 1' if line.startswith('## ') else 'Normal'
 text=line[2:] if style=='Title' else line[3:] if style=='Heading 1' else line
 p=d.add_paragraph(text,style)
 if style!='Normal':p.paragraph_format.keep_with_next=True
 if text.startswith(('Plaasch','Figueres','Plagborg','Giacomini','Corradi','Eurostat.','Chicago Fed.')):
  for r in p.runs:r.font.size=Pt(10)
p=s.footer.paragraphs[0];p.alignment=1;r=p.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
for st in d.styles:
 if st.element.rPr is not None:
  rf=st.element.rPr.find(qn('w:rFonts'))
  if rf is not None:
   for key in list(rf.attrib):
    if 'Theme' in key:del rf.attrib[key]
   for attr in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+attr),'Times New Roman')
 for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for p in d.paragraphs:
 for border in list(p._p.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.save(P/'THESIS_DECISION.docx')
