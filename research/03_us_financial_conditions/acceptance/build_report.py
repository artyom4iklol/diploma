from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml.ns import qn
p=Path(__file__).resolve().parent
d=Document();s=d.sections[0];s.page_width=Inches(8.27);s.page_height=Inches(11.69);s.top_margin=s.bottom_margin=Inches(.8);s.left_margin=s.right_margin=Inches(.85)
for name in ['Normal','Title','Heading 1']:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(6)
d.styles['Normal'].paragraph_format.line_spacing=1.08
d.styles['Title'].font.size=Pt(18);d.styles['Title'].font.bold=True
d.styles['Heading 1'].font.size=Pt(14);d.styles['Heading 1'].paragraph_format.space_before=Pt(12)
for b in (p/'ACCEPTANCE_REPORT.md').read_text().split('\n\n'):
 if b.startswith('# '):d.add_paragraph(b[2:],style='Title')
 elif b.startswith('## '):d.add_paragraph(b[3:],style='Heading 1')
 elif b.strip():
  para=d.add_paragraph(b.strip())
  if b.startswith('[') or 'https://' in b:
   for r in para.runs:r.font.size=Pt(10)
for el in list(d.styles.element.iter()):
 if el.tag==qn('w:rFonts'):
  for a in list(el.attrib):
   if 'theme' in a.lower():del el.attrib[a]
  for a in ['ascii','hAnsi','eastAsia','cs']:el.set(qn('w:'+a),'Times New Roman')
 if el.tag==qn('w:pBdr'):el.getparent().remove(el)
d.core_properties.title='Приемлемость темы диплома';d.core_properties.author='';d.save(p/'ACCEPTANCE_REPORT.docx')
