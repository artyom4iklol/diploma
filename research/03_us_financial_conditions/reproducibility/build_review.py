from pathlib import Path
import json,re,shutil,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
task=Path(__file__).resolve().parent
base=task.parent if task.name=='reproducibility' else task/'deliverables/thesis_us_financial_conditions_10'
reg=json.loads((base/'catalog/registry.json').read_text())
raw=(task/'cards.txt').read_text();cards=[]
parts=re.split(r'(?m)^(S\d\d)\n',raw)
for i in range(1,len(parts),2):
 id=parts[i];fields=[line.split('|',1) for line in parts[i+1].strip().splitlines()];cards.append({'id':id,'fields':fields})
assert len(cards)==34 and {x['id'] for x in cards}=={x['id'] for x in reg}
missing={'S05','S09','S10','S19','S22','S32'}
for x in reg:
 x['review_stage']='10';x['review_date']='2026-10-02'
 x['read_status']='Предварительный разбор доступных аннотаций; полный текст отсутствует' if x['id'] in missing else 'Содержательный разбор по доступной полной версии; без независимой репликации и проверки всех доказательств'
 x['review_card']='review/LITERATURE_REVIEW_ALL_34.md#'+x['id'].lower()
 if x['id']=='S32':x['doi']='10.1016/j.jfs.2022.101004'
 if x['id'] not in missing:x['full_text']['pages']=len(PdfReader(base/x['full_text']['file']).pages)
 if x['id']=='S16':x['version_caution']='В arXiv v3 обнаружены несогласованные даты выборки и размещение частот в таблице 1; численные выводы требуют сверки.'
(base/'catalog/registry.json').write_text(json.dumps(reg,ensure_ascii=False,indent=2))
(base/'review/review_cards.json').write_text(json.dumps(cards,ensure_ascii=False,indent=2))
intro=(task/'synthesis.md').read_text();out=[intro,'\n# Карточки всех работ\n']
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(12)
 st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
 st.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
doc.styles['Heading 2'].font.size=Pt(13);doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
foot=sec.footer.paragraphs[0];foot.alignment=2
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
for block in intro.split('\n\n'):
 if block.startswith('# '):doc.add_paragraph(block[2:].strip(),style='Title')
 elif block.startswith('## '):doc.add_paragraph(block[3:].strip(),style='Heading 1')
 elif block.strip():doc.add_paragraph(block.strip())
doc.add_page_break();doc.add_heading('Карточки всех работ',level=1)
by={r['id']:r for r in reg}
for c in cards:
 x=by[c['id']]
 # Keep bibliography title as body text, allowing scientific punctuation.
 doc.add_heading(c['id']+' '+('Предварительный разбор' if c['id'] in missing else 'Разбор доступной полной версии'),level=2)
 b=f"{x['author'].replace(' and ', '; ')}. {x['title']}. {x.get('journal','')}, {x['year']}."
 p=doc.add_paragraph(b);p.paragraph_format.keep_with_next=True;p.runs[0].bold=True
 link='https://doi.org/'+x['doi'] if x.get('doi') else x['url']
 p=doc.add_paragraph(link);p.paragraph_format.keep_with_next=True
 p.runs[0].font.size=Pt(10)
 out += ['\n## '+c['id']+'\n',b+'\n',link+'\n','Статус: '+x['read_status']+'\n']
 for label,txt in c['fields']:
  p=doc.add_paragraph();p.add_run(label+'. ').bold=True;p.add_run(txt)
  out.append('**'+label+'.** '+txt+'\n')
 # No artificial page break per short card.
md='\n'.join(out)
(base/'review/LITERATURE_REVIEW_ALL_34.md').write_text(md)
doc.core_properties.title='Финансовые условия США в прогнозировании рисков роста еврозоны'
doc.core_properties.subject='Обзор 34 научных работ и уточнение вклада бакалаврского диплома'
doc.core_properties.author=''
# Remove template theme font overrides and decorative paragraph borders.
for el in list(doc.styles.element.iter()):
 if el.tag==qn('w:rFonts'):
  for attr in list(el.attrib):
   if 'theme' in attr.lower():del el.attrib[attr]
  for attr in ['ascii','hAnsi','eastAsia','cs']:el.set(qn('w:'+attr),'Times New Roman')
 if el.tag==qn('w:pBdr'):el.getparent().remove(el)
for el in doc.element.iter():
 if el.tag==qn('w:pBdr'):el.getparent().remove(el)
for run in doc.paragraphs[0].runs:
 run.font.name='Times New Roman'
 run._element.get_or_add_rPr().rFonts.set(qn('w:cs'),'Times New Roman')
doc.save(base/'review/LITERATURE_REVIEW_ALL_34.docx')
# Preserve previous research and provenance without doubling the papers.
old=Path('stage10/previous/thesis_us_financial_conditions_09')
for p in (old.iterdir() if old.exists() else []):
 if p.name in ['papers','catalog','previous_stage']:continue
 dest=base/'history/stage09'/p.name;dest.parent.mkdir(parents=True,exist_ok=True)
 if p.is_dir():shutil.copytree(p,dest,dirs_exist_ok=True)
 else:shutil.copy2(p,dest)
for name in ['access_retry.json','recovery_audit.json','recover_zip.py','retry.py','cards.txt','synthesis.md','build_review.py']:
 dest=base/('review' if name.endswith('.json') else 'reproducibility')/name;dest.parent.mkdir(exist_ok=True);
 if (task/name).resolve()!=dest.resolve() and (task/name).exists():shutil.copy2(task/name,dest)
print(json.dumps({'cards':len(cards),'full_pdf':len(reg)-len(missing),'words':len(md.split()),'docx':str(base/'review/LITERATURE_REVIEW_ALL_34.docx')},ensure_ascii=False))
