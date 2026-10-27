"""Reusable Word format based on the user's approved diploma, without copying its text."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from lxml import etree

SPEC=json.loads((Path(__file__).parent/'format_spec.json').read_text())
ALIGN={'left':WD_ALIGN_PARAGRAPH.LEFT,'right':WD_ALIGN_PARAGRAPH.RIGHT,'center':WD_ALIGN_PARAGRAPH.CENTER,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}

def set_font(font,size=12,bold=False,italic=False):
    font.name='Times New Roman';font.size=Pt(size);font.bold=bold;font.italic=italic;font.color.rgb=RGBColor(0,0,0)
    rpr=font._element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for key in list(rf.attrib):
        if key.endswith('Theme'):del rf.attrib[key]
    for name in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+name),'Times New Roman')
    color=rpr.find(qn('w:color'))
    if color is not None:
        for key in list(color.attrib):
            if key!=qn('w:val'):del color.attrib[key]

def apply_style(doc,name,data,outline=None):
    s=doc.styles[name]if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    s.base_style=doc.styles['Normal']if name!='Normal'else None
    set_font(s.font,data.get('size_pt',12),data.get('bold',False),data.get('italic',False))
    p=s.paragraph_format;p.alignment=ALIGN[data.get('alignment','justify')]
    p.first_line_indent=Cm(data.get('first_line_cm',0));p.left_indent=Cm(0);p.right_indent=Cm(0)
    p.line_spacing=data.get('line_spacing',1.5);p.space_before=Pt(data.get('before_pt',0));p.space_after=Pt(data.get('after_pt',0))
    p.widow_control=True;p.keep_with_next=outline is not None;p.keep_together=outline is not None
    p.page_break_before=bool(data.get('new_page',False))
    if outline is not None:
        pr=s.element.get_or_add_pPr();old=pr.find(qn('w:outlineLvl'))
        if old is not None:pr.remove(old)
        ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),str(outline));pr.append(ol)
    return s

def new_document(*,cover=False):
    d=Document();page=SPEC['page'];s=d.sections[0]
    for name,val in [('page_width',page['width_cm']),('page_height',page['height_cm']),('left_margin',page['left_cm']),('right_margin',page['right_cm']),('top_margin',page['top_cm']),('bottom_margin',page['bottom_cm']),('header_distance',page['header_cm']),('footer_distance',page['footer_cm'])]:setattr(s,name,Cm(val))
    for style in d.styles:
        if hasattr(style,'font'):style.font.color.rgb=RGBColor(0,0,0)
        for color in style.element.iter(qn('w:color')):
            color.attrib.clear();color.set(qn('w:val'),'000000')
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    body=SPEC['body'];apply_style(d,'Normal',{**body,'before_pt':body['space_before_pt'],'after_pt':body['space_after_pt']})
    apply_style(d,'Title',{'size_pt':14,'bold':True,'alignment':'center','first_line_cm':0,'line_spacing':1.5,'after_pt':12})
    apply_style(d,'Thesis Chapter',SPEC['chapter'],outline=0)
    apply_style(d,'Thesis Section',SPEC['section'],outline=1)
    for name,key in [('Thesis Figure','figure_caption'),('Thesis Source','source_note'),('Thesis Table Number','table_number'),('Thesis Table Title','table_title')]:apply_style(d,name,SPEC[key])
    apply_style(d,'Footnote Text',SPEC['footnote'])
    if 'Footnote Reference' not in d.styles:
        d.styles.add_style('Footnote Reference',WD_STYLE_TYPE.CHARACTER)
    set_font(d.styles['Footnote Reference'].font,10)
    d.styles['Footnote Reference'].font.superscript=True
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run=OxmlElement('w:r');prop=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');prop.append(sz);run.append(prop);txt=OxmlElement('w:t');txt.text='1';run.append(txt);field.append(run);p._p.append(field)
    s.different_first_page_header_footer=cover
    d.core_properties.title='Основа дипломного исследования автомобильной торговли Китая и ЕС'
    d.core_properties.subject='Подготовка к написанию диплома по проверенным данным 2021–2025'
    d.core_properties.author='Автор исследования';d.core_properties.keywords='Comext, Китай, ЕС, BEV, PHEV, HEV'
    return d

def add_footnote(doc,paragraph,text):
    """Create genuine footnotes, with no active URL relationship."""
    rel=next((r for r in doc.part.rels.values()if r.reltype==RT.FOOTNOTES),None)
    if rel:
        part=rel.target_part;root=etree.fromstring(part.blob)
    else:
        root=OxmlElement('w:footnotes')
        for ident,typ,tag in [(-1,'separator','separator'),(0,'continuationSeparator','continuationSeparator')]:
            n=OxmlElement('w:footnote');n.set(qn('w:id'),str(ident));n.set(qn('w:type'),typ)
            p=OxmlElement('w:p');r=OxmlElement('w:r');r.append(OxmlElement('w:'+tag));p.append(r);n.append(p);root.append(n)
        part=Part(PackURI('/word/footnotes.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',b'',doc.part.package)
        doc.part.relate_to(part,RT.FOOTNOTES)
    ident=max(int(n.get(qn('w:id')))for n in root)+1
    n=OxmlElement('w:footnote');n.set(qn('w:id'),str(ident));p=OxmlElement('w:p');pr=OxmlElement('w:pPr');ps=OxmlElement('w:pStyle');ps.set(qn('w:val'),'FootnoteText');pr.append(ps);p.append(pr)
    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');rs=OxmlElement('w:rStyle');rs.set(qn('w:val'),'FootnoteReference');rp.append(rs);r.append(rp);r.append(OxmlElement('w:footnoteRef'));p.append(r)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text=' '+text;r.append(t);p.append(r);n.append(p);root.append(n)
    part._blob=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    r=paragraph.add_run();r.style=doc.styles['Footnote Reference'];ref=OxmlElement('w:footnoteReference');ref.set(qn('w:id'),str(ident));r._r.append(ref)
    return ident

def add_toc(doc):
    p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\z');p._p.append(field)
    return p
