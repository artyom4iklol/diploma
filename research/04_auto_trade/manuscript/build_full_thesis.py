"""Build a single thesis DOCX with native equations and native footnotes."""
from pathlib import Path
from copy import deepcopy
import hashlib
import json
import re
import subprocess
import sys
import io
import os
import tempfile
from zipfile import ZipFile
from lxml import etree
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
TMP = Path(os.environ.get('AUTO_THESIS_BUILD_DIR', str(Path(tempfile.gettempdir()) / 'auto_trade_full_thesis')))
TMP.mkdir(parents=True, exist_ok=True)
sys.path.insert(0, str(ROOT / 'writing'))
from thesis_format import new_document, apply_style, set_font, SPEC
from prepare_full_thesis import main as prepare, TITLE


def format_run(run, size=12):
    bold, italic = run.bold, run.italic
    set_font(run.font, size, bold=bool(bold), italic=bool(italic))
    run.bold = bold; run.italic = italic


def style_paragraph(p, size=12, align=None, first_line=1.25, spacing=1.5):
    p.paragraph_format.first_line_indent = Cm(first_line)
    p.paragraph_format.line_spacing = spacing
    p.paragraph_format.widow_control = True
    if align is not None: p.alignment = align
    for r in p.runs: format_run(r, size)


def append_field_run(p, typ, text=None):
    r = OxmlElement('w:r')
    if typ == 'instr':
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = text
    else:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), typ)
    r.append(e); p._p.append(r)



def move_frontmatter(d, body_first, headings, page_map):
    elements = []
    def add(text='', style=None, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, size=12):
        p = d.add_paragraph(text, style=style)
        style_paragraph(p, size, align, 0, 1.3)
        p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
        elements.append(p._p); return p
    add('Национальный исследовательский университет')
    add('«Высшая школа экономики»')
    add('Факультет мировой экономики и мировой политики', before=10)
    add('Образовательная программа «Мировая экономика»')
    add('ВЫПУСКНАЯ КВАЛИФИКАЦИОННАЯ РАБОТА', before=66, after=18)
    add(TITLE, style='Title', size=14, after=32)
    add('Выполнил: Артём Сухаверхов', align=WD_ALIGN_PARAGRAPH.RIGHT, before=30)
    add('Москва, 2026', before=155)
    p = add('Содержание', style='Title', align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=12, size=14)
    p.paragraph_format.page_break_before = True
    first_toc = None
    for index, item in enumerate(headings):
        p = add('', align=WD_ALIGN_PARAGRAPH.LEFT)
        if first_toc is None:
            first_toc = p
            append_field_run(p, 'begin'); append_field_run(p, 'instr', ' TOC \\o "1-2" \\z '); append_field_run(p, 'separate')
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.left_indent = Cm(.6 if item['level'] == 2 else 0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.right_indent = Cm(.65)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17.3), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.add_run(item['text'] + '\t' + str(page_map.get(item['text'], '1')))
        for run in p.runs: format_run(run, 12)
        if index == len(headings) - 1: append_field_run(p, 'end')
    for element in elements: body_first.addprevious(element)


def table_widths(t):
    texts = [c.text for c in t.rows[0].cells]
    n = len(texts)
    if texts[0] == 'Показатель' and n == 3: return [4.0, 5.2, 8.3]
    if texts[0] == 'Год' and n == 5: return [1.3, 2.6, 3.4, 4.2, 6.0]
    if n == 2: return [6.0, 11.5]
    if n == 3: return [5.2, 6.0, 6.3]
    if n == 4:
        if texts[0].startswith('Рассматриваемый'): return [4.1, 4.5, 4.6, 4.3]
        if 'Ряд' in texts[0]: return [8.2, 3.1, 3.1, 3.1]
        return [3.5, 4.7, 4.8, 4.5]
    if n == 5:
        if 'Исход' == texts[0] and texts[1] == 'Категория': return [3.0, 2.7, 3.8, 3.1, 4.9]
        if 'Исход' == texts[0]: return [2.7, 4.7, 2.7, 3.0, 4.4]
        if 'Категория' == texts[0] and 'Квартал' == texts[1]: return [2.6, 2.5, 4.1, 4.1, 4.2]
        if texts[0] == 'Адресат': return [3.4, 3.5, 3.8, 3.4, 3.4]
        return [3.0, 3.4, 3.4, 3.8, 3.9]
    if n == 6:
        if 'Категория' == texts[0] and 'Сравнения' == texts[1]: return [2.6, 2.6, 2.8, 2.7, 3.4, 3.4]
        if 'α' in texts[1]: return [2.6, 1.5, 1.5, 4.1, 4.3, 3.5]
        return [2.1, 2.7, 3.0, 3.3, 3.2, 3.2]
    return [17.5 / n] * n


def build():
    prepare()
    dref = new_document(cover=True)
    apply_style(dref, 'Heading 1', SPEC['chapter'], outline=0)
    apply_style(dref, 'Heading 2', SPEC['section'], outline=1)
    apply_style(dref, 'Heading 3', {**SPEC['section'], 'before_pt': 10, 'after_pt': 5}, outline=2)
    apply_style(dref, 'Thesis Summary', {**SPEC['section'], 'before_pt': 14, 'after_pt': 8})
    dref.save(TMP / 'reference.docx')
    md = (HERE / 'full_thesis_content.md').read_text()
    md = re.sub(r'<!--.*?-->', '', md, flags=re.S)
    md = re.sub(r'(!\[[^\]]*\]\([^)]*)\.svg\)', r'\1.png)', md)
    md = re.sub(r'\]\(((?:chapter[23]_figures|thesis_figures)/[^)]+)\)', lambda m: '](' + str(HERE / m.group(1)) + ')', md)
    def equation(match):
        value = match.group(1)
        tag = re.search(r'\\tag\{([^}]+)\}', value)
        value = re.sub(r'\\tag\{[^}]+\}', '', value).strip()
        if tag: value += r'\qquad\text{(' + tag.group(1) + ')}'
        return '\\[\n' + value + '\n\\]'
    md = re.sub(r'\\\[(.*?)\\\]', equation, md, flags=re.S)
    (TMP / 'pandoc_input.md').write_text(md, encoding='utf8')
    subprocess.run(['pandoc', str(TMP / 'pandoc_input.md'), '-f', 'markdown+tex_math_single_backslash-autolink_bare_uris-implicit_figures',
                    '--reference-doc', str(TMP / 'reference.docx'), '-o', str(TMP / 'base.docx')], check=True)
    d = Document(TMP / 'base.docx')
    for sec in d.sections:
        for name, val in [('page_width', 21), ('page_height', 29.7), ('left_margin', 2.5), ('right_margin', 1), ('top_margin', 2), ('bottom_margin', 2), ('header_distance', 1.25), ('footer_distance', 1.25)]: setattr(sec, name, Cm(val))
    d.sections[0].different_first_page_header_footer = True
    for name, data, outline in [('Normal', {**SPEC['body'], 'before_pt': 0, 'after_pt': 0}, None), ('Heading 1', SPEC['chapter'], 0), ('Heading 2', SPEC['section'], 1), ('Heading 3', SPEC['section'], 2)]: apply_style(d, name, data, outline)
    apply_style(d, 'Thesis Summary', SPEC['section'])
    for style in d.styles:
        if hasattr(style, 'font'):
            style.font.color.rgb = RGBColor(0, 0, 0)
        for color in style.element.iter(qn('w:color')):
            color.attrib.clear(); color.set(qn('w:val'), '000000')
        for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    in_bibliography = False
    for p in d.paragraphs:
        t = p.text.strip()
        if p.style.name.startswith('Heading'):
            if p.style.name == 'Heading 1': in_bibliography = t == 'Список использованных источников'
            if t.startswith('Выводы по '):
                p.style = d.styles['Thesis Summary']
                p.paragraph_format.keep_with_next = True
            elif re.match(r'^[АБВ]\.[1-9]', t): p.style = d.styles['Heading 3']
            for run in p.runs: format_run(run, 14 if p.style.name == 'Heading 1' else 12)
            continue
        if p._p.xpath('.//w:drawing'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(7); p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            continue
        if re.match(r'^(?:Рисунок|Продолжение рисунка) [\dАБВ]+\.\d+\s*[—–-]', t):
            p.style = d.styles['Thesis Figure']; p.paragraph_format.keep_with_next = True
            for run in p.runs: format_run(run, 12)
            continue
        if t.startswith('Источник:'):
            p.style = d.styles['Thesis Source']
            p.paragraph_format.keep_together = True
            for run in p.runs: set_font(run.font, 10, italic=True)
            continue
        if re.match(r'^Таблица [\dАБВ]+\.', t):
            prefix, _, title = t.partition(' — ')
            p.clear(); p.add_run(prefix); p.style = d.styles['Thesis Table Number']; p.paragraph_format.keep_with_next = True
            if prefix == 'Таблица А.7': p.paragraph_format.page_break_before = True
            new = d.add_paragraph(title, style='Thesis Table Title'); p._p.addnext(new._p); new.paragraph_format.keep_with_next = True
            for run in p.runs + new.runs: format_run(run, 10)
            continue
        if t.startswith('Панель А.') or t.startswith('Панель Б.'):
            p.paragraph_format.keep_with_next = True
            style_paragraph(p, 12, WD_ALIGN_PARAGRAPH.CENTER, 0, 1.15)
            continue
        if p._p.xpath('.//m:oMathPara'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.line_spacing = 1
            p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_together = True
            continue
        p.style = d.styles['Normal']
        style_paragraph(p)
        if in_bibliography:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_together = True
            for run in p.runs:
                run.text = re.sub(r'https?://[^\s]+', lambda m: re.sub(r'([/_.?=&-])', lambda x: x.group(1) + '\u200b', m.group(0)), run.text)
        if p._p.xpath('./w:pPr/w:numPr'):
            p.paragraph_format.left_indent = Cm(.7); p.paragraph_format.first_line_indent = Cm(-.7)
    for t in d.tables:
        t.autofit = False
        widths = table_widths(t)
        for col, width in zip(t.columns, widths): col.width = Cm(width)
        pr = t._tbl.tblPr
        layout = pr.find(qn('w:tblLayout'))
        if layout is None: layout = OxmlElement('w:tblLayout'); pr.append(layout)
        layout.set(qn('w:type'), 'fixed')
        width = pr.find(qn('w:tblW')); width.set(qn('w:type'), 'dxa'); width.set(qn('w:w'), str(round(17.5 / 2.54 * 1440)))
        borders = OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'000000'); borders.append(e)
        old = pr.find(qn('w:tblBorders'))
        if old is not None: pr.remove(old)
        pr.append(borders)
        for ri, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr()
            for h in list(trpr.findall(qn('w:trHeight'))): trpr.remove(h)
            cant = OxmlElement('w:cantSplit'); trpr.append(cant)
            if ri == 0:
                repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
            for ci, cell in enumerate(row.cells):
                cell.width = Cm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr(); margins = OxmlElement('w:tcMar')
                for edge, v in [('top',65),('bottom',65),('left',75),('right',75)]:
                    e = OxmlElement('w:'+edge); e.set(qn('w:w'),str(v)); e.set(qn('w:type'),'dxa'); margins.append(e)
                tcpr.append(margins)
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'),'FFFFFF'); tcpr.append(shade)
                for p in cell.paragraphs:
                    p.style = d.styles['Normal']; p.paragraph_format.first_line_indent = Cm(0); p.paragraph_format.left_indent = Cm(0); p.paragraph_format.right_indent = Cm(0)
                    p.paragraph_format.line_spacing = 1; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.keep_with_next = ri == 0 or (len(t.rows) <= 5 and ri < len(t.rows)-1)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 or len(p.text) < 35 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs: set_font(run.font, 12, bold=ri == 0)
    for shape in d.inline_shapes:
        next_p = shape._inline.xpath('ancestor::w:p')[0].getnext()
        caption = ''.join(next_p.itertext()) if next_p is not None else ''
        width_limit = 13.0 if 'Последовательность совместного разбора' in caption else (15.0 if 'Общая и технологическая концентрация' in caption else (15.5 if 'Предлагаемое распределение действий' in caption else 17.0))
        scale = min(Cm(width_limit) / shape.width, Cm(19.2) / shape.height)
        shape.width = round(shape.width * scale); shape.height = round(shape.height * scale)
    # Native footnotes inherit the approved 10 pt single-spacing style.
    rel = next(x for x in d.part.rels.values() if x.reltype.endswith('/footnotes'))
    foot = etree.fromstring(rel.target_part.blob)
    for p in foot.iter(qn('w:p')):
        pr = p.find(qn('w:pPr'))
        if pr is None: pr = OxmlElement('w:pPr'); p.insert(0,pr)
        spacing = pr.find(qn('w:spacing'))
        if spacing is None: spacing = OxmlElement('w:spacing'); pr.append(spacing)
        spacing.set(qn('w:line'),'240'); spacing.set(qn('w:lineRule'),'auto'); spacing.set(qn('w:after'),'0')
        for r in p.iter(qn('w:r')):
            rp = r.find(qn('w:rPr'))
            if rp is None: rp = OxmlElement('w:rPr'); r.insert(0,rp)
            for tag,val in [('sz','20'),('szCs','20')]:
                e=rp.find(qn('w:'+tag))
                if e is None:e=OxmlElement('w:'+tag);rp.append(e)
                e.set(qn('w:val'),val)
            color=OxmlElement('w:color');color.set(qn('w:val'),'000000');rp.append(color)
    rel.target_part._blob = etree.tostring(foot,xml_declaration=True,encoding='UTF-8',standalone=True)
    # Plain visible URLs are retained; active hyperlinks are removed everywhere.
    for part in d.part.package.parts:
        if part.content_type.endswith('xml') and hasattr(part, '_element'):
            root=part._element
            for h in list(root.iter(qn('w:hyperlink'))):
                parent=h.getparent();pos=parent.index(h)
                for child in list(h):parent.insert(pos,child);pos+=1
                parent.remove(h)
    for rId, relation in list(d.part.rels.items()):
        if relation.reltype.endswith('/hyperlink'): del d.part.rels[rId]
    # Separate adjacent native footnote references so 41 and 42 cannot read as 4142.
    for p in d.paragraphs:
        refs=p._p.xpath('.//w:footnoteReference')
        for ref in refs[:-1]:
            run=ref.getparent();following=run.getnext()
            if following is not None and following.tag==qn('w:r') and following.find(qn('w:footnoteReference')) is not None:
                comma=OxmlElement('w:r');pr=OxmlElement('w:rPr')
                align=OxmlElement('w:vertAlign');align.set(qn('w:val'),'superscript');pr.append(align)
                size=OxmlElement('w:sz');size.set(qn('w:val'),'20');pr.append(size);comma.append(pr)
                text=OxmlElement('w:t');text.text=',';comma.append(text);run.addnext(comma)
    headings=[]
    for p in d.paragraphs:
        if p.style.name in ['Heading 1','Heading 2']:
            headings.append({'text':p.text,'level':1 if p.style.name=='Heading 1' else 2})
    page_map_path=TMP/'page_map.json'
    if not page_map_path.exists(): page_map_path=HERE/'full_thesis_page_map.json'
    page_map=json.loads(page_map_path.read_text()) if page_map_path.exists() else {}
    first=d.element.body[0]
    move_frontmatter(d,first,headings,page_map)
    d.core_properties.title=TITLE;d.core_properties.subject='Дипломное исследование автомобильной торговли Китая и ЕС 2021–2025'
    d.core_properties.author='Артём Сухаверхов';d.core_properties.last_modified_by='Артём Сухаверхов'
    path=ROOT/'documents/diploma_auto_trade.docx'
    buffer=io.BytesIO();d.save(buffer);payload=buffer.getvalue()
    with ZipFile(io.BytesIO(payload)) as verify: assert verify.testzip() is None
    temporary=path.with_suffix('.building')
    with temporary.open('wb') as output:
        output.write(payload);output.flush();os.fsync(output.fileno())
    os.replace(temporary,path)
    assert path.stat().st_size==len(payload)
    (TMP/'headings.json').write_text(json.dumps(headings,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    with ZipFile(path) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
        facts={'native_math_objects':len(xml.findall('.//m:oMath',ns)),'native_display_equations':len(xml.findall('.//m:oMathPara',ns)),
               'native_footnote_references':len(xml.findall('.//w:footnoteReference',ns)),'images':len(xml.findall('.//w:drawing',ns)),
               'active_hyperlinks':len(xml.findall('.//w:hyperlink',ns)),'headings':len(headings),'tables':len(d.tables)}
    facts.update({'path':str(path),'file_bytes':len(payload),'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'visual_review':'pending'})
    (HERE/'full_thesis_document_checks.json').write_text(json.dumps(facts,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
    print(json.dumps(facts,ensure_ascii=False))


if __name__=='__main__':build()
